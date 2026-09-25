"""End-to-end checks for the portal and the bot's background work, against a
throwaway database.

    python -m tools.test_portal

What it proves, and how it stays honest:

  * Real sign-in. Every request goes through /login and a real session
    cookie — no dependency overrides — so a scoping bug can't hide behind a
    faked user.
  * Discord is faked, and *loudly*: the portal's REST calls, Discord's
    sign-in (OAuth) and the bot's discord.py objects are all stand-ins, and
    any call that reaches the real network layer fails the run.
  * Isolation is checked from both directions: what a company manager can
    reach must 404, and — because a leak renders as a perfectly healthy page
    with someone else's data on it — every page a manager can open is scanned
    for the other company's names.
  * The important protections are broken on purpose inside the suite and the
    relevant check must go red (see the "broken on purpose" checks).

It never touches the real database (DB_PATH points at a temp file), Discord, or
a mail server.
"""
import asyncio
import os
import re
import shutil
import sys
import tempfile
import time
import types
import warnings
from collections import Counter
from datetime import timedelta
from pathlib import Path

warnings.filterwarnings("ignore")

# Configuration must be in place before config.py is imported: its values are
# read from the environment once, at import.
_TMP = Path(tempfile.mkdtemp(prefix="rms-test-"))
os.environ.update(
    DB_PATH=str(_TMP / "test.sqlite3"),
    DISCORD_BOT_TOKEN="",                      # any accidental real call fails with 401
    PORTAL_SECRET_KEY="test-secret",
    STORAGE_ROOT=str(_TMP / "sites"),
    DISCORD_CLIENT_ID="424242424242",
    DISCORD_CLIENT_SECRET="test-client-secret",
    PORTAL_BASE_URL="https://portal.test",
    SMTP_HOST="smtp.test", SMTP_PORT="587", SMTP_USER="rms@portal.test", SMTP_PASSWORD="pw",
    SMTP_FROM="RMS <rms@portal.test>", SMTP_SECURITY="starttls",
    JOIN_LINK_DAYS="14", SITES_PER_SERVER="450",
)

import discord  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from config import CONFIG  # noqa: E402
from portal import assign, discord_api, insights, inventory, links, servers  # noqa: E402
from portal.auth import hash_password  # noqa: E402
from portal.main import app  # noqa: E402
from storage import db  # noqa: E402
from utils import access, blockers, oauth  # noqa: E402
from utils.phone import normalize, phone_key  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent

results: list[bool] = []
skipped: list[str] = []

# The customer's files. Kept out of git (names and phone numbers), so a CI
# checkout doesn't have them: the checks that need them are skipped, and
# counted, rather than faked.
WORKBOOK = ROOT / "reference" / "site_master_list.xlsx"
CHAT = ROOT / "reference" / "site_chat_sample.txt"
NO_REFERENCE = "the customer's reference files aren't in this checkout"


def section(title: str) -> None:
    print(f"\n{title}")


def check(label: str, ok, detail: str = "") -> bool:
    ok = bool(ok)
    results.append(ok)
    print(f"  {'PASS' if ok else 'FAIL'}  {label}" + (f"   [{detail}]" if detail and not ok else ""))
    return ok


def skip(label: str, why: str) -> None:
    skipped.append(label)
    print(f"  SKIP  {label}   [{why}]")


def run(coro):
    return asyncio.run(coro)


# --- the portal's view of Discord, in memory -----------------------------------

class FakeDiscord:
    """Just enough of Discord's REST API to observe what the portal asked of it."""

    def __init__(self):
        self.reset()

    def reset(self):
        self.overwrites: dict[int, dict[int, tuple[int, int]]] = {}   # channel -> user -> (allow, deny)
        self.roles_removed: list[tuple[int, int]] = []
        self.left: list[int] = []
        self.calls: Counter = Counter()
        self.fail: set[str] = set()

    def _maybe_fail(self, name, status=500):
        self.calls[name] += 1
        if name in self.fail:
            raise discord_api.DiscordError(f"{name} failed on purpose", status=status)


fake = FakeDiscord()
REAL: dict = {}     # the real functions, kept so their own handling can be tested with the network stubbed


def install_fakes():
    for name in ("_request", "set_overwrite", "channel_member_overwrites", "leave_guild", "revoke_channel_access"):
        REAL[name] = getattr(discord_api, name)
    for name in ("exchange_code", "refresh", "get_me"):
        REAL["oauth." + name] = getattr(oauth, name)

    async def refuse(*args, **kwargs):
        raise AssertionError(f"a test reached the real Discord network layer: {args}")

    discord_api._request = refuse

    async def set_overwrite(channel_id, user_id, mode, reason=None):
        fake._maybe_fail("set_overwrite")
        fake.overwrites.setdefault(channel_id, {})[user_id] = access.overwrite_for(mode)

    async def revoke_channel_access(channel_id, user_id, reason=""):
        fake._maybe_fail("revoke_channel_access")
        fake.overwrites.setdefault(channel_id, {}).pop(user_id, None)

    async def channel_member_overwrites(channel_id):
        return dict(fake.overwrites.get(channel_id, {}))

    async def guild_roles(guild_id):
        return {CONFIG.technician_role_name: 555}

    async def remove_member_role(guild_id, user_id, role_id, reason=""):
        fake.roles_removed.append((guild_id, user_id))

    async def leave_guild(guild_id):
        fake._maybe_fail("leave_guild")
        fake.left.append(guild_id)

    for fn in (set_overwrite, revoke_channel_access, channel_member_overwrites, guild_roles,
               remove_member_role, leave_guild):
        setattr(discord_api, fn.__name__, fn)

    # Discord's sign-in: codes and tokens this suite hands out itself.
    async def exchange_code(code):
        signin.calls["exchange"] += 1
        if code not in signin.codes:
            raise oauth.OAuthError("unknown code", invalid_grant=True, status=400)
        return dict(signin.codes.pop(code))

    async def refresh(refresh_token):
        signin.calls["refresh"] += 1
        answer = signin.refreshes.get(refresh_token)
        if isinstance(answer, Exception):
            raise answer
        if answer is None:
            raise oauth.OAuthError("unknown refresh token", invalid_grant=True, status=400)
        return dict(answer)

    async def get_me(access_token):
        signin.calls["me"] += 1
        if access_token not in signin.people:
            raise oauth.OAuthError("unknown token", status=401)
        return dict(signin.people[access_token])

    oauth.exchange_code, oauth.refresh, oauth.get_me = exchange_code, refresh, get_me

    # Mail: a server that only writes down what it was given.
    import smtplib

    class FakeSMTP:
        def __init__(self, host, port, timeout=None, context=None):
            mail.connections.append((host, port))
            if mail.down:
                raise OSError("connection refused")

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def starttls(self, context=None):
            mail.tls += 1

        def login(self, user, password):
            mail.logins.append(user)

        def send_message(self, message):
            mail.sent.append(message)

    smtplib.SMTP = FakeSMTP
    smtplib.SMTP_SSL = FakeSMTP


class SignIn:
    def __init__(self):
        self.codes, self.refreshes, self.people, self.calls = {}, {}, {}, Counter()

    def person(self, code, discord_id, username, access="acc", refresh="ref", scope="identify guilds.join"):
        """Makes `code` sign in as this Discord account."""
        token = f"{access}-{discord_id}-{code}"
        self.people[token] = {"id": str(discord_id), "username": username}
        self.codes[code] = {"access_token": token, "refresh_token": f"{refresh}-{discord_id}-{code}",
                            "expires_in": 604800, "scope": scope, "token_type": "Bearer"}
        return token

    def bot(self, code, guild_id, name="Fresh server"):
        """Makes `code` the answer to adding the bot to this server."""
        self.codes[code] = {"access_token": "bot-add", "refresh_token": "x", "expires_in": 604800,
                            "scope": "bot", "guild": {"id": str(guild_id), "name": name}}


class Mail:
    def __init__(self):
        self.reset()

    def reset(self):
        self.sent, self.logins, self.connections, self.tls, self.down = [], [], [], 0, False


signin = SignIn()
mail = Mail()


# --- fixtures ----------------------------------------------------------------

PASSWORD = "correct-horse-battery"
CH = {"ISB1": 11, "ISB2": 12, "ISB3": 13, "ISB9": 19, "KHI1": 21}
GUILD = {"ISB1": 1001, "ISB2": 1001, "ISB3": 1001, "ISB9": 1001, "KHI1": 1002}
ids: dict = {}


def site_row(site_id, city, **extra):
    row = {c: "" for c in db.SITE_COLUMNS}
    row.update(site_id=site_id, city=city, region="North", rms_category="Indoor + Multiple Rectifiers",
               existing_equipment="Rectifier: HW-ICC500, Huawei; iDMU; Battery: Li x2",
               install_boq="AC Energy Meter (3V3I) x1; RSMS Gateway x1")
    row.update(extra)       # after the defaults, so a caller can override any of them
    return row


def build_fixtures():
    db.init_db()
    ids["A"] = db.ensure_company("HQ Tech", is_internal=True)
    ids["B"] = db.ensure_company("Karachi Tech")

    db.upsert_sites([
        site_row("ISB1", "Islamabad"), site_row("ISB2", "Islamabad"), site_row("ISB3", "Islamabad"),
        site_row("ISB9", "Islamabad"),
        site_row("KHI1", "Karachi", region="South"), site_row("KHI2", "Karachi", region="South"),
        site_row("FREE1", "Islamabad"),
    ], "fixture")
    db.bulk_assign_company(["ISB1", "ISB2", "ISB3"], ids["A"])
    db.bulk_assign_company(["ISB9", "KHI1", "KHI2"], ids["B"])

    for site_id, channel in CH.items():
        city = "Islamabad" if site_id.startswith("ISB") else "Karachi"
        db.upsert_site_channel(site_id, city, GUILD[site_id], channel, f"{site_id[:3]} Sites 001", None,
                               f"**{site_id}** brief text")
    db.claim_guild(1001, "HQ Tech_North", ids["A"], "North", 1, "fixture", status="ready")
    db.claim_guild(1002, "Karachi Tech_South", ids["B"], "South", 1, "fixture", status="ready")

    for key, email, role, company in (
        ("admin", "admin@t.test", "hq_admin", None),
        ("staff", "staff@t.test", "hq_staff", None),
        ("mgrA", "mgra@t.test", "company_manager", ids["A"]),
        ("mgrB", "mgrb@t.test", "company_manager", ids["B"]),
        ("orphan", "orphan@t.test", "company_manager", None),
    ):
        ids[key] = db.create_portal_user(email, hash_password(PASSWORD), key, role,
                                         must_change=False, company_id=company)

    ids["techA"] = db.create_technician(ids["A"], "Alpha Tech", "0300 1111111", None, "fixture")
    ids["techB"] = db.create_technician(ids["B"], "Bravo Tech", "0321 2222222", None, "fixture")

    # Distinctive text, so a leak is visible on the page.
    db.log_message(9101, CH["ISB1"], "ISB1", 1, "Alpha Tech", "ALPHA-CONVERSATION-MARKER", 1.0)
    db.log_message(9102, CH["KHI1"], "KHI1", 2, "Bravo Tech", "BRAVO-CONVERSATION-MARKER", 1.0)
    ids["blockA"] = db.open_blocker("ISB1", CH["ISB1"], 9001, 1, "Alpha Tech", "ALPHA-BLOCKER not working", "not working")
    ids["blockB"] = db.open_blocker("KHI1", CH["KHI1"], 9002, 2, "Bravo Tech", "BRAVO-BLOCKER not working", "not working")

    root = Path(CONFIG.storage_root)
    for key, site_id, msg in (("attA", "ISB1", 7001), ("attB", "KHI1", 7002)):
        folder = root / site_id / "2026-09-23"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"untagged_{msg}_photo.jpg"
        path.write_bytes(b"\xff\xd8\xff not really a jpeg")
        db.log_attachment(msg, CH[site_id], site_id, 1, "photo.jpg", str(path), None, False)
        ids[key] = next(a["id"] for a in db.attachments_for_site(site_id) if a["message_id"] == msg)
        ids[key + "_msg"] = msg


def login(who: str) -> TestClient:
    email = {"admin": "admin@t.test", "staff": "staff@t.test", "mgrA": "mgra@t.test",
             "mgrB": "mgrb@t.test", "orphan": "orphan@t.test"}[who]
    client = TestClient(app, base_url="https://testserver")
    r = client.post("/login", data={"email": email, "password": PASSWORD}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/", f"login failed for {who}: {r.status_code}"
    return client


def anonymous() -> TestClient:
    return TestClient(app, base_url="https://testserver")


def get(client, path, **kw):
    return client.get(path, follow_redirects=False, **kw)


def post(client, path, data=None, **kw):
    return client.post(path, data=data or {}, follow_redirects=False, **kw)


def msg_of(response) -> str:
    """The flash (msg or err) a redirect carries, decoded."""
    from urllib.parse import parse_qs, urlparse
    q = parse_qs(urlparse(response.headers.get("location", "")).query)
    return " ".join(q.get("msg", []) + q.get("err", []))


def word(text: str, token: str) -> bool:
    return re.search(rf"\b{re.escape(token)}\b", text) is not None


_PHONES = iter(range(1000000, 9999999))


def fresh_tech(company, name, phone=None, email=None):
    phone = phone or f"0301 {next(_PHONES):07d}"
    return db.get_technician(db.create_technician(company, name, phone, email, "test"))


def connect(tech, discord_id, username=None):
    """The technician taps their link and connects Discord as this account
    (what the OAuth callback does, minus Discord)."""
    link = db.ensure_link(technician_id=tech["id"], created_by="test")
    problem = db.bind_link(link["id"], discord_id, username or f"user{discord_id}", f"acc-{discord_id}",
                           f"ref-{discord_id}", 604800, "identify guilds.join")
    assert problem is None, problem
    return db.get_technician(tech["id"])


def held_by(site_id):
    return {t["user_id"] for t in db.techs_for_site(site_id)}


def queued_for(tech):
    return {r["site_id"] for r in db.queue_for_technician(tech["id"])}


def expect_error(fn, *args):
    try:
        result = fn(*args)
        if asyncio.iscoroutine(result):
            run(result)
    except assign.AssignError as exc:
        return str(exc)
    return None


# =============================================================================
# 1. TENANCY: what a company manager can and cannot reach
# =============================================================================

def test_tenancy():
    section("1. Tenancy — a company manager reaches only their own company")
    admin, staff = login("admin"), login("staff")
    mgrA, mgrB, orphan = login("mgrA"), login("mgrB"), login("orphan")

    # --- site pages
    check("A sees their own site", get(mgrA, "/sites/ISB1").status_code == 200)
    check("A cannot open B's site (404, not 403)", get(mgrA, "/sites/KHI1").status_code == 404)
    check("B cannot open A's site", get(mgrB, "/sites/ISB1").status_code == 404)
    check("an unallocated site is invisible to a company", get(mgrA, "/sites/FREE1").status_code == 404)
    check("a site that doesn't exist looks identical to one that isn't yours",
          get(mgrA, "/sites/NOPE1").status_code == get(mgrA, "/sites/KHI1").status_code == 404)
    check("HQ admin can open any site", get(admin, "/sites/KHI1").status_code == 200)
    check("HQ staff can open any site", get(staff, "/sites/ISB1").status_code == 200)
    check("a company account with no company sees nothing (fails closed)",
          get(orphan, "/sites/ISB1").status_code == 404 and get(orphan, "/sites/KHI1").status_code == 404)

    # --- lists and dashboards never contain the other side
    def leak_scan(client, own_site, own_tech, other_company, forbidden, label):
        """Opens every page this company can reach and looks for the OTHER
        company's names. `forbidden` is what must never appear."""
        pages = ["/", "/sites", "/technicians", "/blockers", f"/sites/{own_site}",
                 f"/technicians/{own_tech}", f"/sites?company_id={other_company}",
                 f"/technicians?company_id={other_company}"]
        bad = []
        for path in pages:
            r = get(client, path)
            if r.status_code != 200:
                bad.append(f"{path} -> {r.status_code}")
                continue
            for token in forbidden:
                if word(r.text, token):
                    bad.append(f"{path} leaked {token!r}")
        check(f"{label}: no page leaks the other company", not bad, "; ".join(bad))

    leak_scan(mgrA, "ISB1", ids["techA"], ids["B"],
              ["Karachi Tech", "Bravo Tech", "BRAVO-CONVERSATION-MARKER", "BRAVO-BLOCKER", "KHI1", "KHI2", "ISB9"],
              "company A")
    leak_scan(mgrB, "KHI1", ids["techB"], ids["A"],
              ["HQ Tech", "Alpha Tech", "ALPHA-CONVERSATION-MARKER", "ALPHA-BLOCKER", "ISB1", "ISB2", "ISB3"],
              "company B")

    # A sees its own things, so the scan above isn't passing on empty pages
    text = get(mgrA, "/").text
    check("A's dashboard does show A's own sites (so the leak scan isn't passing on an empty page)",
          word(text, "ISB1") and word(text, "ISB2"))
    check("A's sites page lists A's sites and not B's",
          word(get(mgrA, "/sites").text, "ISB1") and not word(get(mgrA, "/sites").text, "KHI1"))
    check("A's dashboard names the company", "HQ Tech" in text)
    page_b = get(mgrB, "/sites").text
    check("a company sees ALL its sites, including one still waiting for a Discord channel",
          word(page_b, "KHI2") and "not in Discord yet" in page_b)
    hq_default = get(admin, "/sites").text
    hq_all = get(admin, "/sites?show=all").text
    check("HQ's Sites page defaults to sites that are in Discord (the inventory lives on its own page)",
          word(hq_default, "KHI1") and not word(hq_default, "KHI2") and not word(hq_default, "FREE1"))
    check("...says how many it left out, and offers them", "in the inventory but not in Discord yet" in hq_default
          and "show=all" in hq_default)
    check("...and 'show all' includes every site", word(hq_all, "KHI2") and word(hq_all, "FREE1"))
    check("company filter on /sites can't be used to peek",
          not word(get(mgrA, f"/sites?company_id={ids['B']}").text, "KHI1"))

    # --- blockers
    page = get(mgrA, "/blockers").text
    check("A's blocker list shows A's blocker only", "ALPHA-BLOCKER" in page and "BRAVO-BLOCKER" not in page)
    r = post(mgrA, f"/blockers/{ids['blockB']}/resolve")
    check("A cannot resolve B's blocker", r.status_code == 404)
    check("...and it is still open", db.get_blocker(ids["blockB"])["status"] == "open")
    r = post(mgrB, f"/blockers/{ids['blockA']}/dismiss")
    check("B cannot dismiss A's blocker", r.status_code == 404 and db.get_blocker(ids["blockA"])["status"] == "open")
    check("empty filter fields don't 422 the blockers page",
          get(mgrA, "/blockers?city=&site_id=&tech=&hours=").status_code == 200)

    # --- progress / dates / completion / assignment on someone else's site
    before = db.get_site_progress("KHI1")
    check("A cannot change B's site progress",
          post(mgrA, "/sites/KHI1/progress", {"state": "on_site"}).status_code == 404
          and db.get_site_progress("KHI1") == before)
    check("A cannot complete or reopen B's survey",
          post(mgrA, "/sites/KHI1/complete").status_code == 404 and post(mgrA, "/sites/KHI1/reopen").status_code == 404
          and db.get_site_progress("KHI1") == before)
    check("A cannot set B's target date",
          post(mgrA, "/sites/KHI1/target-date", {"target_install_date": "2026-12-01"}).status_code == 404
          and (db.get_site("KHI1")["target_install_date"] or "") == "")
    check("A cannot assign into B's site",
          post(mgrA, "/sites/KHI1/assign", {"technician_id": str(ids["techB"])}).status_code == 404)
    check("A cannot revoke on B's site", post(mgrA, "/sites/KHI1/revoke", {"user_id": "5"}).status_code == 404)
    db.queue_sites(ids["techB"], ["KHI2"], "fixture")
    b_row = db.queued_row(ids["techB"], "KHI2")
    check("A cannot cancel B's assignments, even naming B's row",
          post(mgrA, "/sites/KHI2/cancel-queued", {"queue_id": str(b_row["id"])}).status_code == 404
          and post(mgrA, "/sites/ISB1/cancel-queued", {"queue_id": str(b_row["id"])}).status_code == 404
          and db.queued_row(ids["techB"], "KHI2") is not None)
    db.cancel_queue_row(b_row["id"])

    # --- technicians
    check("A cannot open B's technician", get(mgrA, f"/technicians/{ids['techB']}").status_code == 404)
    check("A cannot give B's technician sites",
          post(mgrA, f"/technicians/{ids['techB']}/assign-sites", {"site_ids": ["ISB1"]}).status_code == 404)
    check("A cannot send B's technician their link",
          post(mgrA, f"/technicians/{ids['techB']}/send-link").status_code == 404)
    check("A cannot unassign B's technician",
          post(mgrA, f"/technicians/{ids['techB']}/unassign-site", {"site_id": "KHI1"}).status_code == 404)
    check("A cannot remove B's technician",
          post(mgrA, f"/technicians/{ids['techB']}/remove").status_code == 404
          and db.get_technician(ids["techB"])["removed_at"] is None)
    r = post(mgrA, f"/technicians/{ids['techA']}/assign-sites", {"site_ids": ["ISB2", "KHI1"]})
    check("a list with one of B's sites in it is refused whole — never a partial success",
          r.status_code == 404 and not queued_for(db.get_technician(ids["techA"])))
    check("A cannot assign B's technician on A's own site",
          post(mgrA, "/sites/ISB1/assign", {"technician_id": str(ids["techB"])}).status_code == 404)

    # creating a technician: a manager's company_id form value is ignored
    post(mgrA, "/technicians", {"name": "Sneaky", "phone": "0333 9990001", "company_id": str(ids["B"])})
    created = db.find_technician_by_phone_key(phone_key("0333 9990001"))
    check("A's new technician lands in A, whatever the form says", created and created["company_id"] == ids["A"])
    r = post(mgrA, "/technicians", {"name": "Dup", "phone": "+92 321 2222222"})
    check("A cannot claim a number B already has (any spelling)",
          "another company" in msg_of(r) and db.find_technician_by_phone_key(phone_key("0321 2222222"))["company_id"] == ids["B"])
    r = post(mgrA, "/sites/ISB1/assign", {"name": "Dup2", "phone": "03212222222"})
    check("...nor by quick-adding on a site page", "another company" in msg_of(r))

    # --- admin-only and HQ-only surfaces
    admin_only = ["/inventory", "/companies", "/discord", "/hermes/rules", "/users", "/reconcile"]
    for path in admin_only:
        codes = {who: get(c, path).status_code for who, c in (("A", mgrA), ("staff", staff), ("admin", admin))}
        check(f"{path}: admin only", codes == {"A": 403, "staff": 403, "admin": 200}, str(codes))
    for path, data in (("/inventory/assign", {"site_ids": ["KHI1"], "company_id": str(ids["A"])}),
                       ("/companies", {"name": "Evil Co"}), ("/hermes/rules", {"kind": "keyword", "value": "x"}),
                       ("/users", {"email": "e@e.e", "name": "E", "password": "longenough1", "role": "hq_admin"}),
                       ("/discord/staff", {"name": "Evil", "discord_role": "Management"}),
                       ("/discord/staff/me", {})):
        check(f"POST {path}: refused for a company manager (and HQ staff)",
              post(mgrA, path, data).status_code == 403 and post(staff, path, data).status_code == 403)
    check("adding the bot to a server is admin-only",
          get(mgrA, f"/discord/add-bot?company_id={ids['A']}&region=North&seq=2").status_code == 403
          and get(staff, f"/discord/add-bot?company_id={ids['A']}&region=North&seq=2").status_code == 403)
    check("technical view is HQ-only",
          get(mgrA, "/sites/ISB1/technical").status_code == 403
          and get(staff, "/sites/ISB1/technical").status_code == 200
          and get(admin, "/sites/ISB1/technical").status_code == 200)
    check("a company manager can't mint an admin, a company or a staff member via the forms",
          db.get_portal_user_by_email("e@e.e") is None and db.get_company_by_name("Evil Co") is None
          and not any(s["name"] == "Evil" for s in db.all_staff()))

    # --- media
    check("A can fetch A's file", get(mgrA, f"/media/file/{ids['attA']}").status_code == 200)
    check("A cannot fetch B's file", get(mgrA, f"/media/file/{ids['attB']}").status_code == 404)
    r = post(mgrA, f"/media/{ids['attB']}/tag", {"tag": "rectifier"})
    check("A cannot re-tag B's photo", r.status_code == 404 and db.get_attachment(ids["attB"])["tag"] is None)
    r = post(mgrA, f"/media/{ids['attA']}/tag",
             {"tag": "rectifier", "message_id": str(ids["attB_msg"]), "site_id": "KHI1"})
    check("a forged message_id/site_id in the form is ignored",
          db.get_attachment(ids["attB"])["tag"] is None and db.get_attachment(ids["attA"])["tag"] == "rectifier",
          str((db.get_attachment(ids["attA"])["tag"], db.get_attachment(ids["attB"])["tag"])))

    # --- unauthenticated
    anon = anonymous()
    for path in ("/", "/sites", "/sites/ISB1", "/technicians", "/blockers", "/inventory", "/companies",
                 "/discord", "/hermes/rules", "/live/version"):
        r = get(anon, path)
        check(f"signed-out {path} -> login", r.status_code == 303 and r.headers["location"] == "/login")

    # --- account creation can't produce an unscoped company manager
    post(admin, "/users", {"email": "cm@t.test", "name": "CM", "password": "longenough1",
                           "role": "company_manager", "company_id": ""})
    check("a company manager without a company is refused", db.get_portal_user_by_email("cm@t.test") is None)
    post(admin, "/users", {"email": "cm2@t.test", "name": "CM2", "password": "longenough1",
                           "role": "hq_staff", "company_id": str(ids["A"])})
    check("an HQ role never keeps a company_id, whatever the form says",
          db.get_portal_user_by_email("cm2@t.test")["company_id"] is None)
    post(admin, "/users", {"email": "cm3@t.test", "name": "CM3", "password": "longenough1",
                           "role": "company_manager", "company_id": str(ids["B"])})
    check("a company manager is created scoped to their company",
          db.get_portal_user_by_email("cm3@t.test")["company_id"] == ids["B"])


# =============================================================================
# 2. ASSIGNMENT: one action, one link, and the bot does the rest
# =============================================================================

def test_assignment():
    from urllib.parse import unquote
    section("2. Assignment — tick sites, press Assign; one link per person, once")
    fake.reset()
    mail.reset()
    A, B = ids["A"], ids["B"]

    # --- someone who has never connected Discord
    nadia = fresh_tech(A, "Nadia Khan")
    o = assign.assign_sites(nadia, ["ISB1", "ISB2"], "tester")
    check("assigning records the sites as waiting for them", set(o.queued) == {"ISB1", "ISB2"}
          and queued_for(nadia) == {"ISB1", "ISB2"})
    check("...makes their one personal link", o.link and o.link["technician_id"] == nadia["id"] and not o.connected)
    check("...and touches nothing in Discord (the bot does that once they connect)",
          not fake.overwrites and not fake.calls)
    again = assign.assign_sites(nadia, ["ISB2", "ISB3"], "tester")
    check("assigning more later rides the SAME link", again.link["token"] == o.link["token"])
    check("...and a site already waiting isn't queued twice", again.queued == ["ISB3"] and again.already == ["ISB2"])

    with db.get_conn() as conn:          # the link runs out before they ever use it
        conn.execute("UPDATE join_links SET expires_at = ? WHERE id = ?", (time.time() - 1, o.link["id"]))
    fresh = assign.assign_sites(fresh_tech(A, "Unrelated"), ["ISB1"], "tester")      # someone else, unaffected
    renewed = db.ensure_link(technician_id=nadia["id"], created_by="t")
    check("an expired, never-used link is replaced by a new one, never handed back",
          renewed["token"] != o.link["token"] and db.get_link(o.link["token"])["revoked_at"] is not None)
    check("...other people's links are untouched by that", db.get_link(fresh.link["token"])["revoked_at"] is None)

    # --- once connected, their link is theirs for good
    nadia = connect(nadia, 5001, "nadia.k")
    check("connecting links their Discord account", nadia["discord_id"] == 5001)
    o2 = assign.assign_sites(nadia, ["ISB1"], "tester")
    check("assigning to someone connected says so (the bot opens it; no new link)",
          o2.connected and o2.link["used_at"] is not None and o2.link["token"] == renewed["token"])
    db.assign_tech("ISB1", 5001, "bot")                 # what the bot does when it opens it
    db.finish_queue_row(db.queued_row(nadia["id"], "ISB1")["id"])
    o3 = assign.assign_sites(db.get_technician(nadia["id"]), ["ISB1"], "tester")
    check("a site they already hold is reported as theirs, not queued again", o3.already == ["ISB1"] and not o3.queued)

    # --- what's refused
    wrong = fresh_tech(B, "Wrongco")
    check("a technician can't be given another company's site",
          "different company" in (expect_error(assign.assign_sites, wrong, ["ISB1"], "t") or ""))
    check("an unallocated site can't be assigned",
          "allocated" in (expect_error(assign.assign_sites, nadia, ["FREE1"], "t") or ""))
    check("nothing ticked is refused", "Tick at least one" in (expect_error(assign.assign_sites, nadia, [], "t") or ""))
    db.set_site_progress("ISB3", "done", "t")
    check("a completed survey takes no new technicians", "complete" in (expect_error(assign.assign_sites, fresh_tech(A, "Late"), ["ISB3"], "t") or ""))
    db.set_site_progress("ISB3", "on_site", "t")
    gone = fresh_tech(A, "Gone Soon")
    run(assign.remove_technician(gone))
    check("a removed technician can't be assigned",
          "removed" in (expect_error(assign.assign_sites, db.get_technician(gone["id"]), ["ISB2"], "t") or ""))
    check("a refused list changes nothing at all",
          not queued_for(wrong) and expect_error(assign.assign_sites, wrong, ["KHI1", "ISB1"], "t") and not queued_for(wrong))

    # --- taking back
    tariq = connect(fresh_tech(A, "Tariq"), 3131)
    assign.assign_sites(tariq, ["ISB2"], "t")
    run(assign.unassign(tariq, "ISB2"))
    check("taking back a site still waiting just withdraws it", "ISB2" not in queued_for(tariq))
    db.assign_tech("ISB2", 3131, "bot")
    fake.overwrites[CH["ISB2"]] = {3131: access.overwrite_for(access.WRITABLE)}
    run(assign.unassign(tariq, "ISB2"))
    check("taking back an open site removes the channel access in Discord, at once",
          3131 not in fake.overwrites[CH["ISB2"]] and 3131 not in held_by("ISB2"))

    # --- removing a technician takes everything back, or nothing
    qasim = connect(fresh_tech(A, "Qasim"), 6001)
    db.assign_tech("ISB1", 6001, "bot")
    fake.overwrites.setdefault(CH["ISB1"], {})[6001] = access.overwrite_for(access.WRITABLE)
    assign.assign_sites(qasim, ["ISB2"], "t")
    fake.fail.add("revoke_channel_access")
    refused = expect_error(assign.remove_technician, qasim)
    check("if Discord refuses to take a site back, the technician is NOT removed",
          refused and db.get_technician(qasim["id"])["removed_at"] is None and 6001 in held_by("ISB1"))
    fake.fail.clear()
    run(assign.remove_technician(qasim))
    check("removal takes the channel back", 6001 not in fake.overwrites[CH["ISB1"]] and 6001 not in held_by("ISB1"))
    check("...withdraws what was waiting, and their link", not queued_for(qasim)
          and db.current_link(technician_id=qasim["id"]) is None)
    check("...forgets their Discord sign-in", db.get_tokens(6001) is None)
    check("...and drops them from the roster (their number is free again)",
          db.get_technician(qasim["id"])["removed_at"] and db.find_technician_by_phone_key(qasim["phone_key"]) is None)

    # --- through the web, as a company manager
    fake.reset()
    mail.reset()
    mgrA = login("mgrA")
    r = post(mgrA, "/technicians", {"name": "Sultan Shb", "phone": "6475550123", "country": "CA",
                                    "email": "sultan@example.test"})
    sultan = db.find_technician_by_phone_key("16475550123")
    check("a technician from another country is stored in international form", sultan and sultan["phone"] == "+16475550123"
          and sultan["country"] == "CA", msg_of(r))
    listing = get(mgrA, "/technicians").text
    check("the Existing list shows the number with its country", "+1 647-555-0123 · Canada" in listing)
    check("...and says their link isn't sent yet", "Not sent yet" in listing)
    r = post(mgrA, "/technicians", {"name": "Wrong Country", "phone": "6475550123", "country": "PK"})
    check("the same digits read as Pakistani are refused (a landline, not WhatsApp)",
          "landline" in msg_of(r), msg_of(r))
    r = post(mgrA, "/technicians", {"name": "Sultan Again", "phone": "+1 (647) 555-0123"})
    check("the same number typed another way is the same person", "already on the roster" in msg_of(r)
          and len([t for t in db.technicians_for_company(A) if t["phone_key"] == "16475550123"]) == 1)

    r = post(mgrA, f"/technicians/{sultan['id']}/assign-sites", {"site_ids": ["ISB1", "ISB2"]})
    check("ticking two sites assigns both", queued_for(sultan) == {"ISB1", "ISB2"} and "2 sites assigned" in msg_of(r), msg_of(r))
    check("...and his link is emailed, at once", len(mail.sent) == 1 and mail.sent[0]["To"] == "sultan@example.test"
          and "on its way to sultan@example.test" in msg_of(r))
    body = mail.sent[0].get_content()
    link = db.current_link(technician_id=sultan["id"])
    check("the email carries his link and names his sites",
          f"https://portal.test/j/{link['token']}" in body and "ISB1" in body and "ISB2" in body
          and "Connect Discord" in mail.sent[0]["Subject"])
    check("...and it's recorded as sent", db.last_email(technician_id=sultan["id"])["status"] == "sent")
    page = get(mgrA, f"/technicians/{sultan['id']}").text
    wa = re.search(r'href="(https://wa\.me/[^"]+)"', page)
    check("his page offers Send on WhatsApp, to his number, with his link in the message",
          wa and wa.group(1).startswith("https://wa.me/16475550123?text=")
          and f"https://portal.test/j/{link['token']}" in unquote(wa.group(1).replace("&amp;", "&")))
    check("...shows the link itself, to copy", f"https://portal.test/j/{link['token']}" in page)
    check("...and says the link is sent and waiting", "Link sent" in page and "waiting for them to connect Discord" in page)

    r = post(mgrA, "/sites/ISB3/assign", {"name": "Walk In", "phone": "0333 5550001"})
    walkin = db.find_technician_by_phone_key(phone_key("0333 5550001"))
    check("quick-add on a site page adds and assigns in one step", walkin and queued_for(walkin) == {"ISB3"}, msg_of(r))
    check("...with no email, the operator is told to use WhatsApp", "WhatsApp" in msg_of(r))
    site_page = get(mgrA, "/sites/ISB3").text
    check("the site page lists him as waiting, with a WhatsApp button for his link",
          "Walk In" in site_page and "Send link on WhatsApp" in site_page and "wa.me/923335550001" in site_page)
    row = db.queued_row(walkin["id"], "ISB3")
    post(mgrA, "/sites/ISB3/cancel-queued", {"queue_id": str(row["id"])})
    check("cancelling from the site page withdraws it", not queued_for(walkin))
    check("...and he's back in the roster to pick", "Walk In" in get(mgrA, "/sites/ISB3").text)
    r = post(mgrA, "/sites/ISB1/assign", {"technician_id": ""})
    check("assigning nobody says so", "Pick a technician" in msg_of(r))
    page = get(mgrA, "/sites/ISB2").text
    check("the roster dropdown offers only the company's technicians", "Alpha Tech" in page and "Bravo Tech" not in page)
    r = post(mgrA, f"/technicians/{sultan['id']}/assign-sites", {"site_ids": ["KHI2"]})
    check("a site outside their company is simply not found", r.status_code == 404)

    # --- the link again
    with db.get_conn() as conn:
        conn.execute("UPDATE join_links SET expires_at = ? WHERE id = ?", (time.time() - 1, link["id"]))
    check("the page says when the link has expired", "Link expired" in get(mgrA, f"/technicians/{sultan['id']}").text)
    mail.reset()
    r = post(mgrA, f"/technicians/{sultan['id']}/send-link")
    newer = db.current_link(technician_id=sultan["id"])
    check("sending again after it expired makes a new link, and emails it",
          newer["token"] != link["token"] and len(mail.sent) == 1 and newer["token"] in mail.sent[0].get_content())

    mail.down = True
    post(mgrA, f"/technicians/{sultan['id']}/send-link")
    check("a mail server that won't answer is recorded as a failure, not an error page",
          db.last_email(technician_id=sultan["id"])["status"] == "failed"
          and "failed" in get(mgrA, f"/technicians/{sultan['id']}").text)
    mail.down = False


# =============================================================================
# 3. THE BOT'S BACKGROUND WORK: servers, channels, people, staff
# =============================================================================

BOT_ID = 777000000000000001
_ids = iter(range(700000, 10**7))


def http_error(status, code):
    return discord.HTTPException(types.SimpleNamespace(status=status, reason="fake"), {"code": code, "message": "fake"})


class FRole:
    def __init__(self, name, position=1, permissions=None, rid=None):
        self.id = rid or next(_ids)
        self.name, self.position = name, position
        self.permissions = permissions or discord.Permissions.none()

    def _key(self):
        return (self.position, -self.id)          # as Discord orders roles: position, then the older one

    def __lt__(self, other):
        return self._key() < other._key()

    def __gt__(self, other):
        return self._key() > other._key()

    def __ge__(self, other):
        return not self < other

    def __le__(self, other):
        return not self > other

    async def edit(self, permissions=None, reason=None):
        self.permissions = permissions


class FMember:
    def __init__(self, uid, roles=()):
        self.id, self.roles = uid, list(roles)

    @property
    def top_role(self):
        return max(self.roles)

    async def add_roles(self, *roles, reason=None):
        for role in roles:
            if role not in self.roles:
                self.roles.append(role)

    async def remove_roles(self, *roles, reason=None):
        self.roles = [r for r in self.roles if r not in roles]


class FMessage:
    def __init__(self, author, content, channel):
        self.author, self.content, self.channel = author, content, channel

    async def pin(self, reason=None):
        self.channel.pinned.append(self)


class FThread:
    def __init__(self, name):
        self.id, self.name, self.sent = next(_ids), name, []

    async def send(self, content):
        self.sent.append(content)


class FCategory:
    type = discord.ChannelType.category

    def __init__(self, guild, name, overwrites=None):
        self.id, self.guild, self.name, self.overwrites, self.channels = next(_ids), guild, name, overwrites, []

    async def delete(self, reason=None):
        self.guild.all.remove(self)


class FChannel:
    def __init__(self, guild, name, kind=discord.ChannelType.text, category=None, overwrites=None, topic=None):
        self.id, self.guild, self.name, self.type = next(_ids), guild, name, kind
        self.category, self.overwrites, self.topic = category, overwrites or {}, topic
        self.sent, self.pinned, self.threads = [], [], []
        if category is not None:
            category.channels.append(self)

    async def send(self, content):
        message = FMessage(self.guild.me, content, self)
        self.sent.append(message)
        return message

    def pins(self, limit=None):
        async def pinned():
            for m in list(self.pinned):
                yield m
        return pinned()

    async def create_thread(self, name, type=None, invitable=True):
        thread = FThread(name)
        self.threads.append(thread)
        return thread

    async def delete(self, reason=None):
        self.guild.all.remove(self)
        if self.category is not None:
            self.category.channels.remove(self)


class FGuild:
    def __init__(self, gid, name="New server", owner_id=11, age_days=0, others=(), stock=True):
        self.id, self.name, self.owner_id = gid, name, owner_id
        self.created_at = discord.utils.utcnow() - timedelta(days=age_days)
        everyone = discord.Permissions.none()
        everyone.update(view_channel=True, send_messages=True, connect=True, speak=True,
                        create_instant_invite=True, mention_everyone=True, add_reactions=True)
        self.default_role = FRole("@everyone", position=0, permissions=everyone, rid=gid)
        self.bot_role = FRole("RMS Bot", position=1)
        self.roles = [self.default_role, self.bot_role]
        self.me = FMember(BOT_ID, [self.default_role, self.bot_role])
        self.members = {BOT_ID: self.me, owner_id: FMember(owner_id, [self.default_role])}
        for uid in others:
            self.members[uid] = FMember(uid, [self.default_role])
        self.all = []
        self.system_channel = None
        if stock:
            text, voice = FCategory(self, "Text Channels"), FCategory(self, "Voice Channels")
            general = FChannel(self, "general", category=text)
            talk = FChannel(self, "General", kind=discord.ChannelType.voice, category=voice)
            self.all += [text, voice, general, talk]
            self.system_channel = general
        self.verification_level = discord.VerificationLevel.low
        self.default_notifications = discord.NotificationLevel.only_mentions
        self.edits = []

    @property
    def channels(self):
        return list(self.all)

    @property
    def categories(self):
        return [c for c in self.all if c.type == discord.ChannelType.category]

    @property
    def text_channels(self):
        return [c for c in self.all if c.type == discord.ChannelType.text]

    @property
    def member_count(self):
        return len(self.members)

    def get_member(self, uid):
        return self.members.get(uid)

    async def fetch_member(self, uid):
        return self.members[uid]

    def get_channel(self, cid):
        return next((c for c in self.all if c.id == cid), None)

    async def create_role(self, name, permissions=None, mentionable=False, reason=None):
        role = FRole(name, position=1, permissions=permissions)
        self.roles.append(role)
        return role

    async def create_category(self, name, overwrites=None, reason=None):
        category = FCategory(self, name, overwrites)
        self.all.append(category)
        return category

    async def create_text_channel(self, name, category=None, overwrites=None, topic=None, reason=None):
        channel = FChannel(self, name, category=category, overwrites=overwrites, topic=topic)
        self.all.append(channel)
        return channel

    async def edit(self, reason=None, **changes):
        self.edits.append(changes)
        for key, value in changes.items():
            setattr(self, key, value)


class FHTTP:
    """The few raw calls the worker makes. Adding a member checks the token is
    the account's current one, as Discord does."""

    def __init__(self, bot):
        self.bot, self.puts, self.fail = bot, [], {}
        self.overwrites: dict[int, dict[int, tuple[int, int]]] = {}

    async def request(self, route, json=None, **kw):
        gid, uid = map(int, re.search(r"/guilds/(\d+)/members/(\d+)", route.url).groups())
        self.puts.append((gid, uid, (json or {}).get("access_token")))
        if "put" in self.fail:
            raise self.fail["put"]
        tokens = db.get_tokens(uid)
        if not tokens or tokens["access_token"] != (json or {}).get("access_token"):
            raise http_error(403, 50025)
        guild = self.bot.guilds_by_id[gid]
        if uid in guild.members:
            return None                             # 204: already in
        guild.members[uid] = FMember(uid, [guild.default_role])
        return {}                                   # 201: added

    async def edit_channel_permissions(self, channel_id, target, allow, deny, type, reason=None):
        if "overwrite" in self.fail:
            raise self.fail["overwrite"]
        self.overwrites.setdefault(channel_id, {})[target] = (int(allow), int(deny))

    async def delete_channel_permissions(self, channel_id, target, reason=None):
        self.overwrites.get(channel_id, {}).pop(target, None)


class FBot:
    def __init__(self, *guilds):
        self.guilds_by_id = {g.id: g for g in guilds}
        self.http = FHTTP(self)

    def add(self, guild):
        self.guilds_by_id[guild.id] = guild

    def get_guild(self, gid):
        return self.guilds_by_id.get(gid)


def test_bot_worker():
    from cogs import servers as worker_mod
    from utils import discord_layout as layout

    section("3. The bot's background work — servers, channels, people, staff")
    with db.get_conn() as conn:          # a clean queue: earlier sections' rows would point at servers not in this fake bot
        conn.execute("UPDATE site_queue SET cancelled_at = ? WHERE opened_at IS NULL AND cancelled_at IS NULL", (time.time(),))

    W = db.ensure_company("Worker Co")
    db.upsert_sites([site_row(f"WK{i}", "Karachi", region="South") for i in range(1, 6)], "t")
    db.bulk_assign_company([f"WK{i}" for i in range(1, 4)], W)
    needed = {s["name"]: s["sites"] for s in servers.needed_servers()}
    check("sites with no server ask for one, named Company_Region", needed.get("Worker Co_South") == 3, str(needed))

    bot = FBot()
    worker = worker_mod.Worker(bot)

    # --- a server the bot finds itself in, unasked
    stranger = FGuild(7001, "Someone's server", age_days=400, others=(1, 2, 3))
    bot.add(stranger)
    db.register_guild_seen(7001, stranger.name, stranger.owner_id)
    run(worker_mod.setup_claimed(worker))
    run(worker_mod.provision(worker))
    check("a server nobody claimed is left completely alone",
          db.get_guild(7001)["status"] == "unrecognised" and len(stranger.all) == 4 and not stranger.edits
          and stranger.default_role.permissions.view_channel)

    # --- a new server, claimed through the portal
    fresh = FGuild(7002, "my new server")
    bot.add(fresh)
    db.claim_guild(7002, fresh.name, W, "South", 1, "admin")
    run(worker_mod.setup_claimed(worker))
    row = db.get_guild(7002)
    perms = fresh.default_role.permissions
    check("a newly created server is set up and marked ready", row["status"] == "ready", str(row))
    check("...Discord's stock #general and voice channel are gone", not fresh.all)
    check("...@everyone can't see channels, join voice, invite people or ping everyone",
          not (perms.view_channel or perms.connect or perms.speak or perms.create_instant_invite or perms.mention_everyone)
          and perms.send_messages and perms.add_reactions)
    check("...the four functional roles exist, below the bot",
          {r.name for r in fresh.roles} >= {"Management", "Engineer", "Coordinator", "Technician"}
          and all(fresh.bot_role > r for r in fresh.roles if r.name in layout.ROLE_NAMES))
    check("...it carries its Company_Region name", fresh.name == "Worker Co_South" and row["name"] == "Worker Co_South")
    check("...anyone can post without a verified email, everyone is notified of everything, no join messages",
          fresh.verification_level == discord.VerificationLevel.none
          and fresh.default_notifications == discord.NotificationLevel.all_messages and fresh.system_channel is None)

    # --- a server that isn't new waits for a person
    old = FGuild(7003, "Old server", age_days=30, others=(5, 6))
    bot.add(old)
    db.claim_guild(7003, old.name, W, "Central A", 1, "admin")
    run(worker_mod.setup_claimed(worker))
    check("a server that isn't brand new is held for an HQ admin to check, untouched",
          db.get_guild(7003)["status"] == "review" and len(old.all) == 4 and old.default_role.permissions.view_channel)
    db.set_guild_status(7003, "confirmed")
    run(worker_mod.setup_claimed(worker))
    check("...once confirmed it's set up — but nothing in it is deleted",
          db.get_guild(7003)["status"] == "ready" and len(old.all) == 4 and not old.default_role.permissions.view_channel)

    # --- a role above the bot: say so, don't half-work
    stuck = FGuild(7004, "Stuck")
    stuck.roles.append(FRole("Management", position=5))
    bot.add(stuck)
    db.claim_guild(7004, stuck.name, W, "North", 1, "admin")
    run(worker_mod.setup_claimed(worker))
    check("a server where a role sits above the bot's is not marked ready, and says what to fix",
          db.get_guild(7004)["status"] == "claimed" and "drag the bot's role above Management" in (db.get_guild(7004)["error"] or ""))
    db.set_guild_status(7004, "left")

    # --- channels for the company's sites in that region
    made = run(worker_mod.provision(worker))
    wk = {s: db.get_site_channel(s) for s in ("WK1", "WK2", "WK3")}
    check("every site of that company and region gets its channel", made == 3 and all(wk.values()), str(made))
    category = fresh.categories[0]
    check("...in a category named for its block of 50", category.name == "Sites 001-050" and len(category.channels) == 3)
    channel = fresh.get_channel(wk["WK1"]["channel_id"])
    check("...named after the site", channel.name == "wk1" and "Karachi" in channel.topic)
    check("...with the site brief pinned by the bot, once",
          len(channel.pinned) == 1 and channel.pinned[0].content.startswith("**WK1**") and wk["WK1"]["brief"])
    check("...and the private access thread", len(channel.threads) == 1 and wk["WK1"]["thread_id"] == channel.threads[0].id)
    everyone_ow = channel.overwrites[fresh.default_role]
    staff_ow = [ow for role, ow in channel.overwrites.items() if getattr(role, "name", "") == "Engineer"]
    check("...hidden from @everyone, open to the staff roles",
          everyone_ow.view_channel is False and staff_ow and staff_ow[0].view_channel and staff_ow[0].send_messages)
    check("sites of other companies, or other regions, don't land in it",
          not db.get_site_channel("WK4") and len(fresh.text_channels) == 3)
    again = run(worker_mod.provision(worker))
    check("running it again creates nothing more", again == 0 and len(fresh.text_channels) == 3
          and all(len(fresh.get_channel(wk[s]["channel_id"]).pinned) == 1 for s in wk))

    # a crash between creating the channel and finishing it
    db.bulk_assign_company(["WK4", "WK5"], W)
    half = FChannel(fresh, "wk4", category=category)                  # created, registered, brief never pinned
    fresh.all.append(half)
    db.upsert_site_channel("WK4", "Karachi", 7002, half.id, category.name)
    orphan = FChannel(fresh, "wk5", category=category)                # created, never registered
    fresh.all.append(orphan)
    run(worker_mod.provision(worker))
    check("after a crash, a half-made channel is finished — one brief, one thread",
          len(half.pinned) == 1 and len(half.threads) == 1 and db.get_site_channel("WK4")["brief"])
    check("...and a channel made but never recorded is reused, not duplicated",
          db.get_site_channel("WK5")["channel_id"] == orphan.id and len([c for c in fresh.text_channels if c.name == "wk5"]) == 1)

    # category arithmetic, directly
    full = FGuild(7005, stock=False)
    block = FCategory(full, "Sites 001-050")
    full.all.append(block)
    for i in range(50):
        full.all.append(FChannel(full, f"x{i}", category=block))
    nxt = run(layout.pick_category(full, {c.id for c in block.channels}, {}))
    check("a full category (50 channels) is left for the next block", nxt.name == "Sites 051-100")
    pilot = FGuild(7006, stock=False)
    legacy = FCategory(pilot, "ISB Sites 001-020")
    pilot.all.append(legacy)
    known = set()
    for i in range(20):
        c = FChannel(pilot, f"isb{i}", category=legacy)
        pilot.all.append(c)
        known.add(c.id)
    check("the pilot's own category (holding our channels) is filled before a new one is made",
          run(layout.pick_category(pilot, known, {})) is legacy)
    crowded = FGuild(7007, stock=False)
    for i in range(CONFIG.channel_count_ceiling):
        crowded.all.append(FChannel(crowded, f"c{i}"))
    try:
        run(layout.pick_category(crowded, set(), {}))
        refused = False
    except layout.LayoutError:
        refused = True
    check("a server at the channel ceiling (500 counts categories) refuses more", refused)

    # --- opening sites for technicians
    tn = connect(fresh_tech(W, "Tech New"), 8101, "tech.new")           # connected, not in the server yet
    tm = connect(fresh_tech(W, "Tech Member"), 8102)                    # already in the server
    fresh.members[8102] = FMember(8102, [fresh.default_role])
    assign.assign_sites(tn, ["WK1", "WK2"], "mgr")
    assign.assign_sites(tm, ["WK1"], "mgr")
    db.set_site_progress("WK2", "done", "mgr")
    opened = run(worker_mod.open_queue(worker))
    ch1, ch2 = fresh.get_channel(wk["WK1"]["channel_id"]), fresh.get_channel(wk["WK2"]["channel_id"])
    check("the bot opens the waiting sites", opened == 3 and not queued_for(tn) and not queued_for(tm), str(opened))
    check("someone not in the server is added with their own sign-in (one call)",
          [p[:2] for p in bot.http.puts] == [(7002, 8101)] and 8101 in fresh.members)
    check("someone already in is not added again", all(uid != 8102 for _, uid, _ in bot.http.puts))
    check("an open survey: they can post", bot.http.overwrites[ch1.id][8101] == access.overwrite_for(access.WRITABLE)
          and bot.http.overwrites[ch1.id][8102] == access.overwrite_for(access.WRITABLE))
    check("a completed survey: read-only from the start, and no greeting",
          bot.http.overwrites[ch2.id][8101] == access.overwrite_for(access.READ_ONLY)
          and not any("<@8101>" in m.content for m in ch2.sent))
    check("they're greeted at the bottom of their channel, with the brief",
          any(m.content.startswith("\U0001f44b <@8101>") and "**WK1**" in m.content for m in ch1.sent))
    check("the portal's records match", {8101, 8102} <= held_by("WK1") and 8101 in held_by("WK2"))
    check("they carry the Technician role",
          any(r.name == "Technician" for r in fresh.members[8101].roles))

    # tokens: refresh when close to running out, and only the bot does it
    tr = connect(fresh_tech(W, "Tech Refresh"), 8103)
    with db.get_conn() as conn:
        conn.execute("UPDATE discord_tokens SET expires_at = ? WHERE discord_id = 8103", (time.time() + 60,))
    signin.refreshes["ref-8103"] = {"access_token": "acc-8103-new", "refresh_token": "ref-8103-new", "expires_in": 604800}
    assign.assign_sites(tr, ["WK3"], "mgr")
    run(worker_mod.open_queue(worker))
    tokens = db.get_tokens(8103)
    check("a sign-in about to expire is refreshed and the new one used",
          8103 in fresh.members and tokens["access_token"] == "acc-8103-new" and bot.http.puts[-1][2] == "acc-8103-new")
    check("...and the rotated refresh token is kept", tokens["refresh_token"] == "ref-8103-new")

    td = connect(fresh_tech(W, "Tech Dead"), 8104)
    with db.get_conn() as conn:
        conn.execute("UPDATE discord_tokens SET expires_at = ? WHERE discord_id = 8104", (time.time() + 60,))
    assign.assign_sites(td, ["WK3"], "mgr")                            # no refresh answer: Discord says invalid_grant
    run(worker_mod.open_queue(worker))
    row = db.queued_row(td["id"], "WK3")
    check("a sign-in Discord no longer accepts marks them 'connect again', and nothing is guessed",
          row and row["error"] == "needs_connect" and db.get_tokens(8104)["dead_at"] and 8104 not in fresh.members)
    # Hours later, past the longest retry wait: only 'connect again' can be holding it back.
    check("...and that row waits for them (it isn't retried, however long it's been)",
          all(r["id"] != row["id"] for r in db.openable_queue_rows(now=time.time() + 7200)))
    check("the technician's page says so", "Needs to connect again" in get(login("admin"), f"/technicians/{td['id']}").text)
    connect(db.get_technician(td["id"]), 8104)                          # they tap their link again
    run(worker_mod.open_queue(worker))
    check("...once they reconnect, it opens", not queued_for(td) and 8104 in fresh.members)

    tf = connect(fresh_tech(W, "Tech Flaky"), 8105)
    with db.get_conn() as conn:
        conn.execute("UPDATE discord_tokens SET expires_at = ? WHERE discord_id = 8105", (time.time() + 60,))
    signin.refreshes["ref-8105"] = oauth.OAuthError("Discord unreachable", invalid_grant=False)
    assign.assign_sites(tf, ["WK3"], "mgr")
    run(worker_mod.open_queue(worker))
    row = db.queued_row(tf["id"], "WK3")
    check("if Discord can't be reached, the row waits and tries again later — the sign-in isn't thrown away",
          row and row["error"] and row["attempts"] == 1 and not db.get_tokens(8105)["dead_at"])

    tx = connect(fresh_tech(W, "Tech Full"), 8106)
    bot.http.fail["put"] = http_error(400, 30001)
    assign.assign_sites(tx, ["WK3"], "mgr")
    run(worker_mod.open_queue(worker))
    bot.http.fail.clear()
    check("Discord's limits are explained in words", "100 Discord servers" in (db.queued_row(tx["id"], "WK3")["error"] or ""))

    # races with the portal
    tc = connect(fresh_tech(W, "Tech Cancel"), 8107)
    fresh.members[8107] = FMember(8107, [fresh.default_role])
    assign.assign_sites(tc, ["WK3"], "mgr")
    pending = next(r for r in db.openable_queue_rows(100) if r["technician_id"] == tc["id"])
    db.claim_queue_row(pending["id"])
    db.cancel_queue_row(pending["id"])                                  # withdrawn while the bot works
    run(worker_mod.open_row(worker, pending))
    ch3 = fresh.get_channel(wk["WK3"]["channel_id"])
    check("withdrawn in the portal while the bot worked: the access is taken back again",
          8107 not in bot.http.overwrites.get(ch3.id, {}) and 8107 not in held_by("WK3"))

    tk = connect(fresh_tech(W, "Tech Complete"), 8108)
    fresh.members[8108] = FMember(8108, [fresh.default_role])
    assign.assign_sites(tk, ["WK3"], "mgr")
    pending = next(r for r in db.openable_queue_rows(100) if r["technician_id"] == tk["id"])
    real_mode, answers = db.site_mode, iter(["writable", "read_only", "read_only"])
    db.site_mode = lambda site_id: next(answers, "read_only")          # completed while the bot worked
    try:
        db.claim_queue_row(pending["id"])
        run(worker_mod.open_row(worker, pending))
    finally:
        db.site_mode = real_mode
    check("completed in the portal while the bot worked: they end up read-only",
          bot.http.overwrites[ch3.id][8108] == access.overwrite_for(access.READ_ONLY))

    tb = connect(fresh_tech(W, "Tech Broken"), 8109)
    assign.assign_sites(tb, ["WK4"], "mgr")
    db.upsert_site_channel("WK4", "Karachi", 7002, 999999999, category.name)    # its channel is gone from Discord
    tg = connect(fresh_tech(W, "Tech Good"), 8110)
    assign.assign_sites(tg, ["WK5"], "mgr")
    run(worker_mod.open_queue(worker))
    check("one row that fails doesn't stop the next",
          "channel no longer exists" in (db.queued_row(tb["id"], "WK4")["error"] or "") and not queued_for(tg))

    # --- HQ staff, in every server
    sid = db.create_staff("Ayesha", "Management", email="ayesha@t.test", created_by="t")
    link = db.ensure_link(staff_id=sid, created_by="t")
    db.bind_link(link["id"], 8201, "ayesha", "acc-8201", "ref-8201", 604800, "identify guilds.join")
    run(worker_mod.sync_staff(worker))
    roles_of = lambda g, uid: {r.name for r in g.members[uid].roles}
    check("a connected staff member is put into every ready server, with their role",
          "Management" in roles_of(fresh, 8201) and "Management" in roles_of(old, 8201))
    check("...and it's recorded", {(r["guild_id"], bool(r["ok"])) for r in db.guild_staff_rows() if r["staff_id"] == sid}
          >= {(7002, True), (7003, True)})
    with db.get_conn() as conn:
        conn.execute("UPDATE staff SET discord_role = 'Engineer' WHERE id = ?", (sid,))
    run(worker_mod.sync_staff(worker))
    check("changing their role swaps it", roles_of(fresh, 8201) >= {"Engineer"} and "Management" not in roles_of(fresh, 8201))
    db.remove_staff(sid)
    run(worker_mod.sync_staff(worker))
    check("a removed staff member loses the staff role everywhere (they stay in the server)",
          not (roles_of(fresh, 8201) & {"Management", "Engineer", "Coordinator"}) and 8201 in fresh.members)

    # --- the loop keeps going when one part fails
    db.upsert_sites([site_row("WK6", "Karachi", region="South")], "t")
    db.bulk_assign_company(["WK6"], W)
    real_open = worker_mod.open_queue

    async def exploding(w):
        raise RuntimeError("boom")
    worker_mod.open_queue = exploding
    try:
        run(worker_mod.tick(worker))
    finally:
        worker_mod.open_queue = real_open
    check("a failing step is logged and the rest of the tick still runs", db.get_site_channel("WK6") is not None)
    check("...including the heartbeat the portal watches", db.bot_online())

    # --- servers coming and going
    cog = worker_mod.ServersCog.__new__(worker_mod.ServersCog)
    run(worker_mod.ServersCog.on_guild_join(cog, FGuild(7008, "Added by hand")))
    check("a server the bot is added to outside the portal is recorded as unrecognised",
          db.get_guild(7008)["status"] == "unrecognised")
    run(worker_mod.ServersCog.on_guild_remove(cog, fresh))
    check("a server the bot is removed from is marked gone", db.get_guild(7002)["status"] == "left")
    admin = login("admin")
    page = get(admin, "/discord").text
    check("the Discord page offers to release its sites", "Release its sites" in page and "Worker Co_South" in page)
    post(admin, "/discord/guilds/7002/release")
    check("releasing puts its sites back to waiting, with their technicians still assigned",
          db.get_site_channel("WK1") is None and queued_for(tn) >= {"WK1", "WK2"} and not held_by("WK1"))
    check("...and the region asks for a new server", any(s["name"].startswith("Worker Co_South")
                                                         for s in servers.needed_servers()))
    check("releasing a server that's still in use is refused",
          post(admin, "/discord/guilds/7003/release").status_code == 404)

    # capture never acts in a server nobody claimed
    stray = FChannel(stranger, "general")
    check("the capture cog only knows registered site channels (so an unclaimed server is inert)",
          db.site_id_for_channel(stray.id) is None)


# =============================================================================
# 4. SITES: the template, uploads by a company and by HQ, and allocation
# =============================================================================

def template_csv(*rows, headings=None) -> bytes:
    """A filled-in template as CSV (the same parser reads .xlsx and .csv)."""
    import csv as _csv
    import io as _io

    buf = _io.StringIO()
    writer = _csv.writer(buf)
    writer.writerow(headings or [h for h, _, _ in inventory.TEMPLATE_COLUMNS])
    for row in rows:
        writer.writerow([row.get(field, "") for _, field, _ in inventory.TEMPLATE_COLUMNS])
    return buf.getvalue().encode()


def upload(client, data, name="sites.csv", company_id=None):
    form = {} if company_id is None else {"company_id": str(company_id)}
    return client.post("/sites/upload", data=form, files={"file": (name, data, "application/octet-stream")},
                       follow_redirects=False)


def test_inventory():
    import io
    import time as _time
    from datetime import datetime as _dt
    from openpyxl import Workbook, load_workbook
    from tools.site_template import write_template

    section("4. Sites — the template, uploads by a company and by HQ, and allocation")
    regions = db.all_regions()
    check("the region list starts with the workbook's four regions", set(db.DEFAULT_REGIONS) <= set(regions), str(regions))

    # --- the template itself
    wb = load_workbook(io.BytesIO(inventory.build_template(regions)))
    ws = wb["Sites"]
    check("the template's headings are exactly its columns, required ones starred",
          [c.value for c in ws[1]] == [h + (" *" if req else "") for h, _, req in inventory.TEMPLATE_COLUMNS])
    rules = ws.data_validations.dataValidation
    check("...Region is a dropdown drawn from the region list",
          any(v.type == "list" and "B2" in str(v.sqref) for v in rules)
          and [c.value for c in wb["Lists"]["A"]] == regions)
    check("...the list sheet is hidden, and there's a 'How to fill' sheet",
          wb["Lists"].sheet_state == "hidden" and "How to fill" in wb.sheetnames)
    check("...and Site ID cells are text, so 0013 stays 0013", ws["A2"].number_format == "@")

    # --- the customer's workbook, through the CLI and back in. Without it (a CI
    # checkout), 1,200 made-up sites take its place so the scale checks still run.
    if WORKBOOK.exists():
        from tools.build_site_csv import build          # the customer-specific converter; not in every checkout
        df, source = build(WORKBOOK), "the customer's workbook"
    else:
        import pandas as pd
        cities = ["Karachi", "Lahore", "Multan", "Taxila", "Sukkur"]
        df = pd.DataFrame([{"site_id": f"SYN{i:04d}", **{c: f"{c} {i}" for c in db.SITE_COLUMNS},
                            "region": regions[i % len(regions)], "city": cities[i % len(cities)]}
                           for i in range(1200)])
        source = "1,200 made-up sites (no workbook in this checkout)"
    full = _TMP / "all_sites.xlsx"
    written = write_template(df, full)
    rows = inventory.parse_sites_file(full.name, full.read_bytes(), regions)
    by_id = {r["site_id"]: r for r in rows}
    same = len(by_id) == len(df) and all(str(by_id[row["site_id"]][c]) == str(row[c])
                                          for _, row in df.iterrows() for c in db.SITE_COLUMNS)
    check(f"{source} -> template (CLI) -> upload reads all 1,200 sites cell for cell",
          written == 1200 and len(rows) == 1200 and same)
    south = _TMP / "south_50.xlsx"
    check("the CLI picks a region and a count (the demo file: 50 South sites)",
          write_template(df, south, region="South", limit=50) == 50
          and {r["region"] for r in inventory.parse_sites_file(south.name, south.read_bytes(), regions)} == {"South"})

    # --- reading, and what's refused (all-or-nothing)
    def err_of(name, data):
        try:
            inventory.parse_sites_file(name, data, regions)
        except inventory.UploadError as exc:
            return str(exc)
        return None

    good = inventory.parse_sites_file("g.csv", template_csv(
        {"site_id": "nw1", "region": "south", "city": "Karachi", "target_install_date": "2026-10-01"},
        {},                                                                  # a blank row is skipped
        {"site_id": "NW2", "region": "North", "city": "Taxila", "tenants": "Jazz, Zong"}), regions)
    check("a filled-in template reads: ids upper-cased, region as listed, date as ISO, blank rows skipped",
          [(r["site_id"], r["region"], r["target_install_date"]) for r in good]
          == [("NW1", "South", "2026-10-01"), ("NW2", "North", None)] and good[1]["tenants"] == "Jazz, Zong")
    shuffled = [h for h, _, _ in reversed(inventory.TEMPLATE_COLUMNS)]
    buf = io.StringIO()
    import csv as _csv
    w = _csv.writer(buf)
    w.writerow(shuffled)
    w.writerow(["2026-11-05" if h.startswith("Planned") else {"Site ID": "NW3", "Region": "Central A", "City": "Lahore"}.get(h, "") for h in shuffled])
    check("column order doesn't matter, only the headings",
          inventory.parse_sites_file("s.csv", buf.getvalue().encode(), regions)[0]["site_id"] == "NW3")
    check("a row with no City is refused, naming the row",
          "row 2: no City" in (err_of("c.csv", template_csv({"site_id": "NW4", "region": "North"})) or ""))
    check("a region that isn't on the list is refused (a typo would ask for a server nobody needs)",
          "Region 'Soth' isn't one of" in (err_of("r.csv", template_csv({"site_id": "NW5", "region": "Soth", "city": "X"})) or ""))
    check("a date that isn't YYYY-MM-DD is refused",
          "isn't YYYY-MM-DD" in (err_of("d.csv", template_csv({"site_id": "NW6", "region": "North", "city": "X",
                                                                "target_install_date": "01/10/2026"})) or ""))
    check("an unusable Site ID is refused", "isn't a usable Site ID" in (err_of("i.csv", template_csv(
        {"site_id": "bad id!", "region": "North", "city": "X"})) or ""))
    check("a site listed twice is refused", "listed more than once: NW7" in (err_of("t.csv", template_csv(
        {"site_id": "NW7", "region": "North", "city": "X"}, {"site_id": "nw7", "region": "North", "city": "Y"})) or ""))
    many_errors = err_of("m.csv", template_csv(*[{"site_id": f"E{i}", "region": "Nowhere", "city": "X"} for i in range(9)]))
    check("several problems: the first few are named, the rest counted, nothing imported",
          many_errors and "and 5 more" in many_errors and "Nothing was imported" in many_errors, str(many_errors))
    headings = [h for h, _, _ in inventory.TEMPLATE_COLUMNS]
    check("a file missing a column isn't the template",
          "missing Region" in (err_of("h.csv", template_csv(headings=[h for h in headings if h != "Region"])) or ""))
    check("...nor is one with an extra column",
          "unexpected Notes" in (err_of("x.csv", template_csv(headings=headings + ["Notes"])) or ""))
    too_many = template_csv(*[{"site_id": f"Z{i:04d}", "region": "North", "city": "X"} for i in range(inventory.MAX_ROWS + 1)])
    check(f"more than {inventory.MAX_ROWS} sites in one file is refused", "more than" in (err_of("big.csv", too_many) or ""))
    check("an oversized file is refused", "10 MB" in (err_of("b.csv", b"x" * (inventory.MAX_UPLOAD_BYTES + 1)) or ""))
    check("an empty file is refused", "empty" in (err_of("e.csv", b"") or ""))
    check("the wrong kind of file is refused", "site template" in (err_of("x.pdf", b"%PDF") or ""))
    check("garbage pretending to be a workbook is refused politely",
          "couldn't be read" in (err_of("g.xlsx", b"not a workbook at all") or ""))

    risky = inventory.build_template(regions, [{"site_id": "FX1", "region": "North", "city": "X",
                                                "tenants": "=HYPERLINK(\"http://x\")"}])
    cell = load_workbook(io.BytesIO(risky))["Sites"]["J2"]
    check("text that looks like a formula is stored as text, never a formula",
          cell.data_type == "s" and inventory.parse_sites_file("f.xlsx", risky, regions)[0]["tenants"].startswith("=HYPERLINK"))
    typed = Workbook()
    typed.active.title = "Sites"
    typed.active.append([h for h, _, _ in inventory.TEMPLATE_COLUMNS])
    typed.active.append(["ISB0013", "North", "Islamabad"] + [""] * (len(inventory.TEMPLATE_COLUMNS) - 4)
                        + [_dt(2026, 10, 3)])
    typed_buf = io.BytesIO()
    typed.save(typed_buf)
    check("a date Excel stored as a date is read as that day",
          inventory.parse_sites_file("t.xlsx", typed_buf.getvalue(), regions)[0]["target_install_date"] == "2026-10-03")

    # --- what an upload does to the inventory
    A, B = ids["A"], ids["B"]
    db.upsert_sites([site_row("SPARE1", "Lahore", region="Central A"), site_row("LEG1", "Islamabad")], "inventory")
    db.upsert_site_channel("LEG1", "Islamabad", 1001, 61, "legacy", None, "brief")     # unallocated, but has a channel
    result = db.apply_site_upload([
        {**site_row("NW1", "Karachi", region="South"), "target_install_date": "2026-10-01"},
        {**site_row("SPARE1", "Lahore", region="Central A"), "target_install_date": None},
        {**site_row("KHI2", "Karachi"), "target_install_date": None},
        {**site_row("LEG1", "Islamabad"), "target_install_date": None},
    ], A, "mgrA", "first.csv")
    check("a company's upload: new sites are added for it, an unallocated one is taken",
          (result["added"], result["claimed"], result["updated"]) == (1, 1, 0)
          and db.get_site("NW1")["company_id"] == A and db.get_site("SPARE1")["company_id"] == A, str(result))
    check("...another company's site is skipped and untouched",
          ("KHI2", "belongs to another company") in result["skipped"] and db.get_site("KHI2")["company_id"] == B)
    check("...and an unallocated site that already has a Discord channel isn't taken",
          any(s == "LEG1" and "Discord channel" in why for s, why in result["skipped"]) and db.get_site("LEG1")["company_id"] is None)
    check("...with the planned date from the file", db.get_site("NW1")["target_install_date"] == "2026-10-01")

    again = db.apply_site_upload([{**{c: "" for c in db.SITE_COLUMNS}, "site_id": "NW1", "region": "South",
                                   "city": "Karachi", "district": "KARACHI EAST", "target_install_date": None}], A, "mgrA")
    after = db.get_site("NW1")
    check("uploading its own site again refreshes it; blank cells keep what was there",
          again["updated"] == 1 and after["district"] == "KARACHI EAST" and after["rms_category"]
          and after["target_install_date"] == "2026-10-01")
    moved = db.apply_site_upload([{**site_row("ISB1", "Islamabad", region="South"), "target_install_date": None}], A, "mgrA")
    check("a site that has a Discord channel keeps its region (a channel can't move servers)",
          moved["skipped"] and "region can't change" in moved["skipped"][0][1] and db.get_site("ISB1")["region"] == "North")
    inv_only = db.apply_site_upload([{**site_row("HQ1", "Multan", region="Central B"), "target_install_date": None},
                                     {**site_row("KHI2", "Karachi", district="KHI SOUTH"), "target_install_date": None}],
                                    None, "admin")
    check("HQ loading inventory only: new sites arrive unallocated, and nobody's site changes owner",
          db.get_site("HQ1")["company_id"] is None and db.get_site("KHI2")["company_id"] == B
          and db.get_site("KHI2")["district"] == "KHI SOUTH", str(inv_only))
    logged = db.recent_site_imports(A)
    check("every upload is logged against its company, with what it did",
          len(logged) >= 3 and all(i["company_id"] == A for i in logged)
          and any(i["filename"] == "first.csv" and (i["added"], i["claimed"], i["skipped"]) == (1, 1, 2) for i in logged))

    # --- allocation: a site with a Discord channel stays with its company
    assigned, skipped = db.bulk_assign_company(["ISB1"], B)
    check("a site with a Discord channel can't be moved to another company",
          assigned == [] and "Discord channel" in skipped[0][1] and db.get_site("ISB1")["company_id"] == A, str(skipped))
    assigned, skipped = db.bulk_assign_company(["ISB1"], None)
    check("...nor taken back to unallocated", assigned == [] and db.get_site("ISB1")["company_id"] == A)
    assigned, _ = db.bulk_assign_company(["ISB1"], A, "2026-11-02")
    check("re-allocating to the same company just updates the date",
          assigned == ["ISB1"] and db.get_site("ISB1")["target_install_date"] == "2026-11-02")
    assigned, skipped = db.bulk_assign_company(["FREE1", "NOPE-SITE"], A)
    check("a site with no channel can be given; an unknown one is reported",
          assigned == ["FREE1"] and skipped == [("NOPE-SITE", "not in the inventory")])
    assigned, _ = db.bulk_assign_company(["FREE1"], None)
    check("...and taken back again", assigned == ["FREE1"] and db.get_site("FREE1")["company_id"] is None)

    # --- through the web
    admin, mgr_a, staff = login("admin"), login("mgrA"), login("staff")
    n = len(db.all_sites())
    check("a company manager can't download the template or upload sites (only HQ adds sites)",
          get(mgr_a, "/sites/template.xlsx").status_code == 403
          and upload(mgr_a, template_csv({"site_id": "MGR1", "region": "South", "city": "Karachi"})).status_code == 403
          and db.get_site("MGR1") is None and len(db.all_sites()) == n)
    check("...nor can HQ staff, or a login with no company",
          get(staff, "/sites/template.xlsx").status_code == 403 and get(login("orphan"), "/sites/template.xlsx").status_code == 403
          and upload(staff, template_csv({"site_id": "WEB2", "region": "North", "city": "X"})).status_code == 403
          and db.get_site("WEB2") is None)
    r = get(admin, "/sites/template.xlsx")
    check("an HQ admin downloads the template",
          r.status_code == 200 and r.headers["content-type"].startswith(inventory.XLSX_MEDIA_TYPE)
          and "attachment" in r.headers.get("content-disposition", ""))
    check("an HQ admin must say which company",
          "Choose which company" in msg_of(upload(admin, template_csv({"site_id": "WEB2", "region": "North", "city": "X"}))))
    r = upload(admin, template_csv({"site_id": "WEB1", "region": "South", "city": "Karachi"},
                                   {"site_id": "KHI1", "region": "South", "city": "Karachi"}), company_id=A)
    check("the admin's upload lands in the company picked; another company's site is left alone",
          r.status_code == 303 and db.get_site("WEB1")["company_id"] == A and db.get_site("KHI1")["company_id"] == B, msg_of(r))
    check("...and says what happened, naming what was skipped",
          "Imported into HQ Tech: 1 new" in msg_of(r) and "KHI1 (belongs to another company)" in msg_of(r), msg_of(r))
    n = len(db.all_sites())
    r = upload(admin, b"junk", name="bad.xlsx", company_id=B)
    check("a bad upload explains itself and changes nothing", "couldn't be read" in msg_of(r) and len(db.all_sites()) == n)
    own_page = get(mgr_a, "/sites").text
    check("a manager's Sites page has no way to add sites",
          "live-imports" not in own_page and "/sites/upload" not in own_page and "template.xlsx" not in own_page
          and "/sites/upload" not in get(mgr_a, "/").text)
    check("...and HQ's Inventory lists every import", "live-imports" in get(admin, "/inventory").text)

    r = upload(admin, full.read_bytes(), name="all_sites.xlsx", company_id="none")
    check("HQ can load the whole inventory unallocated from the template", "1200 new" in msg_of(r), msg_of(r))
    started = _time.time()
    r = upload(admin, full.read_bytes(), name="all_sites.xlsx", company_id="none")
    check("...and again, which only refreshes (in a sensible time)",
          "0 new" in msg_of(r) and "1200 updated" in msg_of(r) and _time.time() - started < 10, msg_of(r))
    many = [row["site_id"] for row in rows[:1000]]
    assigned, _ = db.bulk_assign_company(many, B)
    check("allocating a thousand sites at once works (past SQLite's variable limit)", len(assigned) == 1000)
    db.bulk_assign_company(many, None)

    picks = [rows[i]["site_id"] for i in (10, 11, 12)]
    r = post(admin, "/inventory/assign", {"site_ids": picks, "company_id": str(B), "target_install_date": "2026-10-20"})
    check("allocating ticked sites through the page", "3 site(s) now allocated to Karachi Tech" in msg_of(r), msg_of(r))
    check("...with the planned date", all(db.get_site(s)["target_install_date"] == "2026-10-20" for s in picks))
    r = post(admin, "/inventory/assign", {"site_ids": ["ISB1"], "company_id": str(B)})
    check("...a site with a channel is left alone with a plain explanation", "left alone" in msg_of(r) and "ISB1" in msg_of(r), msg_of(r))
    check("no sites ticked is refused", "Tick at least one" in msg_of(post(admin, "/inventory/assign", {"company_id": str(B)})))
    check("no company chosen is refused", "Choose a company" in msg_of(post(admin, "/inventory/assign", {"site_ids": picks})))
    check("an impossible date is refused, and nothing moves",
          "valid" in msg_of(post(admin, "/inventory/assign", {"site_ids": picks, "company_id": "none", "target_install_date": "2026-13-45"}))
          and db.get_site(picks[0])["company_id"] == B)
    post(admin, "/inventory/assign", {"site_ids": picks, "company_id": "none"})
    check("'take back' returns them to unallocated", all(db.get_site(s)["company_id"] is None for s in picks))

    # --- regions
    r = post(admin, "/companies/regions", {"name": "  Gilgit   Baltistan "})
    check("HQ adds a region (tidied), and it joins the template's dropdown",
          "Gilgit Baltistan" in db.all_regions() and "Added region" in msg_of(r)
          and "Gilgit Baltistan" in [c.value for c in load_workbook(io.BytesIO(get(admin, "/sites/template.xlsx").content))["Lists"]["A"]])
    check("...the same region in other case isn't added twice",
          "already a region" in msg_of(post(admin, "/companies/regions", {"name": "gilgit baltistan"})))
    check("only an HQ admin adds regions", post(mgr_a, "/companies/regions", {"name": "X"}).status_code == 403
          and post(staff, "/companies/regions", {"name": "X"}).status_code == 403)

    # --- a site that isn't in Discord yet still has a page
    r = get(mgr_a, "/sites/WEB1")
    check("a company's new site (no channel yet) has a page that says it's waiting for its server",
          r.status_code == 200 and "waiting for its server" in r.text and "isn't in Discord yet" in r.text)
    check("...and the lists link to it", 'href="/sites/WEB1"' in get(mgr_a, "/sites").text)

    # --- it stays usable at full size
    started = _time.time()
    pages = {p: get(admin, p) for p in ("/", "/sites", "/inventory", "/companies")}
    took = _time.time() - started
    check("dashboard, sites, inventory and companies all render with 1,200+ sites",
          all(r.status_code == 200 for r in pages.values()), str({p: r.status_code for p, r in pages.items()}))
    check("...in a sensible time", took < 6.0, f"{took:.1f}s for four pages")
    check("the inventory page can filter", get(admin, "/inventory?company=none&provisioned=no&city=").status_code == 200)


# =============================================================================
# 5. THE NUMBERS: forecast, ranking, technical view
# =============================================================================

def test_insights():
    from datetime import date
    section("5. Numbers — forecast, best performer, technical view")

    def row(state="not_started", planned=None, company=1):
        return {"state": state, "target_install_date": planned, "company_id": company}

    rows = [
        row(planned="2026-09-22"),                 # yesterday -> overdue
        row(planned="2026-09-23"),                 # today -> this week
        row(planned="2026-09-27"),                 # Sunday -> this week
        row(planned="2026-09-28"),                 # next Monday -> next week
        row(planned="2026-10-01"),                 # a Thursday: next week AND next month
        row(planned="2026-10-15"),                 # next month only
        row(planned="2026-10-04"),                 # next Sunday -> next week AND next month
        row(planned="2026-11-15"),                 # further out -> nowhere
        row(state="done", planned="2026-09-22"),   # finished sites never count
        row(planned=None),                         # allocated, no date
        row(planned=None, company=None),           # not allocated: nothing to schedule
        row(planned="not-a-date"),                 # garbage counts as no date, not a crash
    ]
    f = insights.forecast(rows, date(2026, 9, 23))
    got = {k: f[k]["count"] for k in ("overdue", "this_week", "next_week", "next_month", "unscheduled")}
    check("forecast buckets a Wednesday correctly",
          got == {"overdue": 1, "this_week": 2, "next_week": 3, "next_month": 3, "unscheduled": 2}, str(got))
    check("...with the date ranges it prints",
          (f["this_week"]["start"], f["this_week"]["end"]) == (date(2026, 9, 23), date(2026, 9, 27))
          and (f["next_week"]["start"], f["next_week"]["end"]) == (date(2026, 9, 28), date(2026, 10, 4))
          and (f["next_month"]["start"], f["next_month"]["end"]) == (date(2026, 10, 1), date(2026, 10, 31)))
    f = insights.forecast([row(planned="2026-09-27")], date(2026, 9, 27))
    check("on a Sunday, 'this week' is just today", f["this_week"]["count"] == 1 and f["next_week"]["count"] == 0)
    f = insights.forecast([row(planned="2027-01-31"), row(planned="2027-02-01")], date(2026, 12, 15))
    check("across a year end, next month is January",
          f["next_month"]["count"] == 1 and f["next_month"]["start"] == date(2027, 1, 1) and f["next_month"]["end"] == date(2027, 1, 31))

    # --- best performer
    co = db.ensure_company("Rank Co")
    db.upsert_sites([{**site_row(f"RK{i}", "Islamabad")} for i in (1, 2, 3, 4)], "test")
    db.bulk_assign_company(["RK1", "RK2", "RK3", "RK4"], co)
    for i in (1, 2, 3, 4):
        db.upsert_site_channel(f"RK{i}", "Islamabad", 1001, 300 + i, "RK", None, None)
    techs = {}
    for name, discord_id in (("Tia", 101), ("Tom", 102), ("Tess", 103), ("Toby", 104), ("Tina", 105)):
        techs[name] = db.create_technician(co, name, f"0345 0000{discord_id}", None, "t")
        db.link_technician_discord_id(techs[name], discord_id)
    for site_id, who in (("RK1", 101), ("RK2", 101), ("RK1", 102), ("RK3", 102), ("RK2", 103), ("RK4", 103), ("RK1", 104)):
        db.assign_tech(site_id, who, "t")
    db.revoke_tech("RK1", 104)                       # Toby finished RK1 then was moved on — it still counts
    for site_id in ("RK1", "RK2"):
        db.set_site_progress(site_id, "done", "t")
    db.open_blocker("RK3", 303, 5001, 102, "Tom", "not working", "not working")   # Tom's still-open site has a blocker
    ranking, best = insights.technician_ranking(co, insights.portfolio(company_id=co))
    order = [r["technician"]["name"] for r in ranking]
    check("best performer is whoever finished the most sites", best and best["technician"]["name"] == "Tia" and best["completed"] == 2, str(order))
    check("ties are broken by fewest open blockers on sites still in progress",
          order.index("Tess") < order.index("Tom"), str(order))
    check("a site finished, then handed on, still counts for the technician who finished it",
          next(r for r in ranking if r["technician"]["name"] == "Toby")["completed"] == 1)
    check("a technician with nothing done ranks last with zero", order[-1] == "Tina" and ranking[-1]["completed"] == 0)
    db.set_site_progress("RK1", "not_started", "t")
    db.set_site_progress("RK2", "not_started", "t")
    _, best = insights.technician_ranking(co, insights.portfolio(company_id=co))
    check("with nothing completed there is no 'best performer' to name", best is None)

    # --- company summaries
    summaries, unassigned = insights.company_summaries(insights.portfolio())
    mine = next(s for s in summaries if s["name"] == "Rank Co")
    check("company summary counts sites, technicians and blockers",
          mine["sites"] == 4 and mine["technicians"] == 5 and mine["open_blockers"] == 1, str(mine))
    check("...and counts what nobody has been given", unassigned["sites"] > 1000)

    # --- the technical view
    db.upsert_sites([{**site_row("TECH1", "Islamabad",
                                 existing_equipment="Rectifier: HW; iDMU; Battery: Li x2; DG1: 40 kVA; SBEE AC meter x2",
                                 install_boq="AC Energy Meter (3V3I) x1; RSMS Gateway x1; DC CT x1")}], "t")
    root = Path(CONFIG.storage_root) / "TECH1"
    root.mkdir(parents=True, exist_ok=True)
    for msg, name, tag, blurry in ((8001, "a.jpg", "rectifier", False), (8002, "b.png", "grid_meter", False),
                                   (8003, "c.pdf", None, False), (8004, "d.ogg", None, False), (8005, "e.jpg", None, True)):
        db.log_attachment(msg, 1, "TECH1", 1, name, str(root / name), tag, blurry)
    v = insights.technical_view("TECH1")
    check("files are counted by kind", (v["photos"], v["documents"], v["media"]) == (3, 1, 1), str((v["photos"], v["documents"], v["media"])))
    check("configurations = distinct kinds of equipment photographed", v["configurations"] == 2)
    check("untagged and blurry are reported", v["untagged"] == 3 and v["blurry"] == 1)
    status = {r["label"]: r["status"] for r in v["reconciliation"]}
    check("a photographed rectifier is marked photographed", status["Rectifier: HW"] == "documented")
    check("a never-photographed iDMU is marked missing", status["iDMU"] == "missing")
    check("an AC meter matches a grid-meter photo", status["SBEE AC meter x2"] == "documented" and status["AC Energy Meter (3V3I) x1"] == "documented")
    check("items with no tag to check are listed but never judged",
          status["Battery: Li x2"] == "unchecked" and status["RSMS Gateway x1"] == "unchecked" and status["DC CT x1"] == "unchecked")
    check("a generator is judged by its own tags", status["DG1: 40 kVA"] == "missing")
    check("the tally counts only judged items", (v["missing"], v["checked"]) == (2, 5), str((v["missing"], v["checked"])))
    check("file kinds classify sensibly",
          [insights.classify_file(n) for n in ("a.JPG", "b.pdf", "c.ogg", "d.mp4", "e.xyz", "")]
          == ["photo", "document", "media", "media", "other", "other"])
    page = get(login("admin"), "/sites/TECH1/technical")
    check("the technical page renders even for a site that isn't in Discord", page.status_code == 200 and "Nothing yet" in page.text)


# =============================================================================
# 6. BLOCKER RULES: the same behaviour, now editable
# =============================================================================

def test_rules():
    from config import BLOCKER_KEYWORDS, BLOCKER_NEGATIONS, BLOCKER_PATTERNS
    section("6. Blocker rules — identical behaviour, editable from the portal")

    seeded = Counter(r["kind"] for r in db.list_hermes_rules())
    check("every built-in rule was seeded, once",
          seeded == Counter({"keyword": len(BLOCKER_KEYWORDS), "negation": len(BLOCKER_NEGATIONS), "pattern": len(BLOCKER_PATTERNS)}), str(seeded))

    compiled = tuple(re.compile(p, re.IGNORECASE) for p in BLOCKER_PATTERNS)

    def legacy(text):
        if not text:
            return None
        lowered = " ".join(text.lower().split())
        for phrase in BLOCKER_NEGATIONS:
            if phrase in lowered:
                return None
        for keyword in BLOCKER_KEYWORDS:
            if keyword in lowered:
                return keyword
        for pattern in compiled:
            m = pattern.search(lowered)
            if m:
                return m.group(0)
        return None

    corpus = []
    if CHAT.exists():
        for line in CHAT.read_text(encoding="utf-8", errors="replace").splitlines():
            body = line.split(" - ", 1)[-1]
            corpus.append(body.split(": ", 1)[-1])
    real = len(corpus)
    corpus += ["no power at site", "koi masla nahi", "still huawei is not readable in putty", "data is not walking in mib",
               "kaam ho gaya done", "bijli nahi hai", "working fine", "", "   ", "Everything OK", "ye kharab hai"]
    blockers.reset_cache()
    differs = [t for t in corpus if blockers.detect(t) != legacy(t)]
    check(f"the database-driven detector agrees with the original on all {len(corpus)} real and probe messages",
          not differs, str(differs[:3]))
    if real:
        flagged = sum(1 for t in corpus if blockers.detect(t))
        check("...and it still flags a sensible minority of the real thread", 0 < flagged < len(corpus) * 0.25, f"{flagged}/{len(corpus)}")
    else:
        skip("...and it still flags a sensible minority of the real thread", NO_REFERENCE)

    # --- editing
    admin = login("admin")
    before = blockers.detect("this is zorrofoo")
    post(admin, "/hermes/rules", {"kind": "keyword", "value": "  ZorroFOO  "})
    check("a keyword added in the portal takes effect (normalised to lower case)",
          before is None and blockers.detect("this is zorrofoo") == "zorrofoo")
    post(admin, "/hermes/rules", {"kind": "negation", "value": "zorrofoo is fine"})
    check("a 'not a blocker' phrase overrides it", blockers.detect("zorrofoo is fine") is None)
    rule = next(r for r in db.list_hermes_rules() if r["kind"] == "keyword" and r["value"] == "masla")
    check("a built-in keyword matches before it's switched off", blockers.detect("masla") == "masla")
    post(admin, f"/hermes/rules/{rule['id']}/toggle")
    check("switching a rule off stops it matching", blockers.detect("masla") is None)
    post(admin, f"/hermes/rules/{rule['id']}/toggle")
    check("...and back on restores it", blockers.detect("masla") == "masla")
    r = post(admin, "/hermes/rules", {"kind": "pattern", "value": "([unclosed"})
    check("an invalid pattern is refused, not stored", "valid pattern" in msg_of(r)
          and not any(x["value"] == "([unclosed" for x in db.list_hermes_rules()))
    post(admin, "/hermes/rules", {"kind": "pattern", "value": r"\bwibble\s+wobble\b"})
    check("a valid pattern works", blockers.detect("the wibble  wobble happened") is not None)
    r = post(admin, "/hermes/rules", {"kind": "keyword", "value": "zorrofoo"})
    check("adding a rule that already exists is a no-op", "already on" in msg_of(r))
    check("an unknown rule type is refused", post(admin, "/hermes/rules", {"kind": "nonsense", "value": "x"}).status_code == 400)

    # The bot is a separate process: it never sees the portal's cache reset, so an edit
    # must reach it through the time-to-live alone, with no restart.
    import time as _time
    real_ttl = blockers.CACHE_SECONDS
    blockers.CACHE_SECONDS = 0.4
    blockers.reset_cache()
    check("(before the edit) the phrase isn't flagged", blockers.detect("quuxfrob broke") is None)
    db.add_hermes_rule("keyword", "quuxfrob", "another process")      # written straight to the database
    check("the old rules are still cached, so it isn't seen instantly", blockers.detect("quuxfrob broke") is None)
    _time.sleep(0.6)
    check("...but is picked up once the cache expires, with no restart and no cache reset",
          blockers.detect("quuxfrob broke") == "quuxfrob")
    blockers.CACHE_SECONDS = real_ttl
    check("the real TTL is about a minute", 30 <= real_ttl <= 120, str(real_ttl))
    blockers.reset_cache()

    with db.get_conn() as conn:      # a bad row that slipped in some other way must not disable detection
        conn.execute("INSERT INTO hermes_rules (kind, value, active, created_by, created_at) VALUES ('pattern', '(((', 1, 'x', 0)")
    blockers.reset_cache()
    check("a corrupt pattern row can't disable detection for everyone", blockers.detect("no power") is not None)

    real_list = db.list_hermes_rules

    def broken(active_only=False):
        raise RuntimeError("database is locked")
    db.list_hermes_rules = broken
    blockers.reset_cache()
    check("if the rules table can't be read, the built-in defaults keep detection alive",
          blockers.detect("no power") is not None and blockers.detect("koi masla nahi") is None)
    db.list_hermes_rules = real_list
    blockers.reset_cache()

    db.init_db()
    check("a restart doesn't resurrect a rule that was switched off",
          (next(r for r in db.list_hermes_rules() if r["id"] == rule["id"])["active"] == 1))
    post(admin, f"/hermes/rules/{rule['id']}/toggle")
    db.init_db()
    check("...(off stays off across restarts)",
          next(r for r in db.list_hermes_rules() if r["id"] == rule["id"])["active"] == 0)
    post(admin, f"/hermes/rules/{rule['id']}/toggle")
    blockers.reset_cache()


# =============================================================================
# 7. THE DATABASE: phone numbers, links, and migrating the live database
# =============================================================================

def test_database():
    import sqlite3
    from utils.phone import display, wa_digits
    section("7. Database — phone numbers, personal links, and migrating the live database")

    same = {normalize(x)[0] for x in ("0300 1234567", "+92 300 1234567", "923001234567", "0092-300-1234567",
                                      "3001234567", "(0300) 123 4567")}
    check("every way of writing one Pakistani number is the same number", same == {"+923001234567"}, str(same))
    check("a number typed with its +code wins over the picked country",
          normalize("+1 647-555-0123", "PK") == ("+16475550123", "16475550123", "CA"))
    check("a Canadian mobile picked as Canada", normalize("647 555 0123", "CA")[0] == "+16475550123")

    def refused(raw, country="PK"):
        try:
            normalize(raw, country)
        except ValueError as exc:
            return str(exc)
        return None
    check("the same digits read as Pakistani are refused: a landline can't be a WhatsApp number",
          "landline" in (refused("6475550123") or ""))
    check("a number that isn't valid for its country is refused", "valid" in (refused("0300 12345") or ""))
    check("text with no digits is refused", "no digits" in (refused("n/a") or ""))
    check("numbers are shown with their country", display("+923001234567") == "+92 300 1234567 · Pakistan"
          and display("+16475550123", "CA") == "+1 647-555-0123 · Canada")
    check("wa.me gets digits only", wa_digits("+92 300-1234567") == "923001234567")
    check("the old key still reads a bare Pakistani number the same way", phone_key("0300 1234567") == "923001234567")

    t1 = db.create_technician(ids["A"], "Dup Test", "0311 0000001", None, "t")
    try:
        db.create_technician(ids["B"], "Dup Test 2", "+92 311 0000001", None, "t")
        dup = False
    except sqlite3.IntegrityError:
        dup = True
    check("a number can't be on two rosters, however it's typed", dup)
    db.remove_technician(t1)
    check("...but it's free again once that technician is removed",
          db.create_technician(ids["B"], "Second", "+92 311 0000001", None, "t") > t1)
    stored = db.get_technician(db.create_technician(ids["A"], "Stored", "0311-0000002", None, "t"))
    check("a technician is stored in international form, with their country",
          stored["phone"] == "+923110000002" and stored["country"] == "PK" and stored["phone_key"] == "923110000002")

    # --- a person's link, and connecting through it
    tech = fresh_tech(ids["A"], "Link Person")
    link = db.ensure_link(technician_id=tech["id"], created_by="t")
    check("a person gets one link, reused until it's used or expires",
          db.ensure_link(technician_id=tech["id"], created_by="t")["token"] == link["token"])
    check("a link token is long and unguessable", len(link["token"]) >= 30)
    bind = lambda link_id, discord_id: db.bind_link(link_id, discord_id, "u", f"a{discord_id}", f"r{discord_id}", 604800, "identify guilds.join")
    with db.get_conn() as conn:
        conn.execute("UPDATE join_links SET expires_at = ? WHERE id = ?", (time.time() - 1, link["id"]))
    check("an expired link can't be used to connect for the first time", bind(link["id"], 9901) == "expired")
    fresh_link = db.ensure_link(technician_id=tech["id"], created_by="t")
    check("...a new one replaces it, and the old one is withdrawn",
          fresh_link["token"] != link["token"] and bind(link["id"], 9901) == "gone")
    check("the new one connects", bind(fresh_link["id"], 9901) is None and db.get_technician(tech["id"])["discord_id"] == 9901
          and db.get_tokens(9901)["access_token"] == "a9901" and db.get_user(9901)["name"] == "Link Person")
    with db.get_conn() as conn:
        conn.execute("UPDATE join_links SET expires_at = ? WHERE id = ?", (time.time() - 1, fresh_link["id"]))
    check("once used, a link stays theirs: the same account can connect through it again, even after its deadline",
          bind(fresh_link["id"], 9901) is None)
    check("...but no other account can", bind(fresh_link["id"], 9902) == "other"
          and db.get_technician(tech["id"])["discord_id"] == 9901)
    other = fresh_tech(ids["A"], "Other Person")
    other_link = db.ensure_link(technician_id=other["id"], created_by="t")
    check("an account already connected for another technician is refused", bind(other_link["id"], 9901) == "taken"
          and db.get_technician(other["id"])["discord_id"] is None)
    staff_id = db.create_staff("Staffer", "Engineer", created_by="t")
    staff_link = db.ensure_link(staff_id=staff_id, created_by="t")
    check("...but the same person can also be HQ staff with the same account", bind(staff_link["id"], 9901) is None)
    db.remove_technician(other["id"])
    check("a removed technician's link is dead", bind(other_link["id"], 9903) == "gone")
    try:
        db.ensure_link(created_by="t")
        both_missing = False
    except ValueError:
        both_missing = True
    check("a link belongs to exactly one person", both_missing)

    # --- a database shaped like the live one before this release
    legacy = _TMP / "legacy.sqlite3"
    conn = sqlite3.connect(legacy)
    conn.executescript("""
        CREATE TABLE portal_users (id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL, name TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'coordinator',
            must_change_password INTEGER NOT NULL DEFAULT 1, created_at REAL NOT NULL, last_login REAL);
        CREATE TABLE pending_invites (token TEXT PRIMARY KEY, invite_code TEXT NOT NULL, guild_id INTEGER NOT NULL,
            channel_id INTEGER NOT NULL, site_id TEXT NOT NULL, name TEXT NOT NULL, phone TEXT NOT NULL,
            created_by TEXT, created_at REAL NOT NULL, consumed_at REAL, discord_id INTEGER, technician_id INTEGER);
        CREATE TABLE technicians (id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER NOT NULL, name TEXT NOT NULL,
            email TEXT, phone TEXT NOT NULL, phone_key TEXT NOT NULL, discord_id INTEGER, created_by TEXT,
            created_at REAL NOT NULL, removed_at REAL);
        CREATE TABLE sites (site_id TEXT PRIMARY KEY, region TEXT, city TEXT, company_id INTEGER, target_install_date TEXT,
            uploaded_by TEXT, created_at REAL NOT NULL, updated_at REAL NOT NULL);
        INSERT INTO portal_users (email, password_hash, name, role, must_change_password, created_at)
            VALUES ('a@x', 'h', 'A', 'admin', 0, 1), ('e@x', 'h', 'E', 'engineer', 0, 1);
        INSERT INTO pending_invites VALUES ('tok', 'CODE', 1, 2, 'ISB1', 'Old', '123', 'me', 1, 5, 7, NULL);
        INSERT INTO technicians (company_id, name, phone, phone_key, created_at, removed_at)
            VALUES (1, 'Removed Sultan', '6475550123', '6475550123', 1, 2);
        INSERT INTO sites VALUES ('ISB1', 'North', 'Islamabad', 1, NULL, 'x', 1, 1), ('GB1', 'Gilgit', 'Gilgit', NULL, NULL, 'x', 1, 1);
    """)
    conn.commit()
    conn.close()

    real_config = db.CONFIG
    db.CONFIG = types.SimpleNamespace(db_path=str(legacy))
    try:
        db.init_db()
        db.init_db()     # the bot and the portal both do this on every start
        c = sqlite3.connect(legacy)
        tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        roles = dict(c.execute("SELECT email, role FROM portal_users").fetchall())
        country = c.execute("SELECT country FROM technicians").fetchone()[0]
        regions = sorted(r[0] for r in c.execute("SELECT name FROM regions"))
        rules = c.execute("SELECT COUNT(*) FROM hermes_rules").fetchone()[0]
        c.close()

        # Two processes starting together: the column appears between the check and the ALTER.
        real_connect = db._connect

        class Racing:
            def __init__(self, conn):
                self.conn = conn

            def execute(self, sql, *args):
                if sql.startswith("PRAGMA table_info(technicians)"):
                    return iter([{"name": "id"}, {"name": "company_id"}])      # "country" looks missing
                return self.conn.execute(sql, *args)

            def __getattr__(self, name):
                return getattr(self.conn, name)

            def __enter__(self):
                self.conn.__enter__()
                return self

            def __exit__(self, *exc):
                return self.conn.__exit__(*exc)

        db._connect = lambda: Racing(real_connect())
        try:
            db.init_db()
            raced = True
        except sqlite3.OperationalError:
            raced = False
        finally:
            db._connect = real_connect
    finally:
        db.CONFIG = real_config

    check("old roles are renamed in place", roles == {"a@x": "hq_admin", "e@x": "hq_staff"}, str(roles))
    check("the retired invite table is gone", "pending_invites" not in tables)
    check("the new tables exist", {"guilds", "staff", "guild_staff", "discord_tokens", "join_links", "site_queue",
                                   "emails", "regions", "site_imports"} <= tables, str(tables))
    check("existing technicians get a country (every one of them was Pakistani)", country == "PK")
    check("the region list is seeded once, with the workbook's regions and any already in the inventory",
          regions == ["Central A", "Central B", "Gilgit", "North", "South"], str(regions))
    from config import BLOCKER_KEYWORDS, BLOCKER_NEGATIONS, BLOCKER_PATTERNS
    expected = len(BLOCKER_KEYWORDS) + len(BLOCKER_NEGATIONS) + len(BLOCKER_PATTERNS)
    check("the rules were seeded, and only once (two starts, one set of rules)", rules == expected, f"{rules} vs {expected}")
    check("two processes adding the same column at once don't crash each other", raced)


# =============================================================================
# 8. SECURITY: things a hostile or careless input must not be able to do
# =============================================================================

def test_security():
    from urllib.parse import unquote
    section("8. Security — hostile input")
    admin = login("admin")

    r = anonymous().get("/healthz")
    check("/healthz answers without signing in, and says nothing else",
          r.status_code == 200 and r.json() == {"ok": True} and r.headers.get("cache-control") == "no-store")
    from tools import healthcheck
    before = db.get_meta("bot_heartbeat", 0)
    db.set_meta("bot_heartbeat", time.time())
    fresh = healthcheck.main("bot")
    db.set_meta("bot_heartbeat", time.time() - 600)
    stale = healthcheck.main("bot")
    db.set_meta("bot_heartbeat", before)
    check("the bot's container health check follows its heartbeat",
          fresh == 0 and stale == 1 and healthcheck.main("nonsense") == 2)

    script = "<script>alert(1)</script>"
    tid = db.create_technician(ids["A"], script, "0322 0000001", None, "t")
    db.create_company('<img src=x onerror=alert(1)>', created_by="t")
    db.create_staff(script, "Coordinator", phone="+923220000002", created_by="t")
    link = db.ensure_link(technician_id=tid, created_by="t")
    db.queue_sites(tid, ["ISB2"], "t")
    pages = [get(admin, p).text for p in ("/technicians", f"/technicians/{tid}", "/sites/ISB2", "/companies",
                                          "/users", "/", "/discord")]
    landing = get(anonymous(), f"/j/{link['token']}").text
    check("a script tag in a name is shown as text, never run — in every page, and on the technician's own link",
          all(script not in p for p in pages + [landing]) and any("&lt;script&gt;" in p for p in pages))
    check("...and so is markup in a company name", all("<img src=x" not in p for p in pages))
    check("no name is ever placed inside a JavaScript string",
          all(script not in p.split("confirm(")[i].split(")")[0] for p in pages for i in range(1, p.count("confirm(") + 1)))
    wa = re.search(r'href="(https://wa\.me/[^"]+)"', get(admin, f"/technicians/{tid}").text)
    check("a name in a WhatsApp message is encoded into the link, never raw HTML",
          wa and "<script>" not in wa.group(1) and "<script>" in unquote(wa.group(1)))

    # redirects after a blocker action
    def resolve_with(referer):
        bid = db.open_blocker("ISB1", CH["ISB1"], 900000 + resolve_with.n, 1, "x", "x not working", "not working")
        resolve_with.n += 1
        headers = {"referer": referer} if referer else {}
        return post(admin, f"/blockers/{bid}/resolve", headers=headers).headers["location"]
    resolve_with.n = 0
    check("resolving goes back to the page you were on", resolve_with("https://testserver/sites/ISB1?msg=x") == "/sites/ISB1?msg=x")
    check("a foreign Referer can't redirect off-site", resolve_with("https://evil.example/steal") == "/steal")
    check("a '//' path (which browsers read as another site) is refused", resolve_with("https://x.example//evil.example/a") == "/blockers")
    check("no Referer falls back to the blocker list", resolve_with(None) == "/blockers")

    # headers
    r = get(anonymous(), f"/j/{link['token']}")
    check("a personal link's page is never cached, never framed, and never passes itself on in a Referer",
          r.headers.get("referrer-policy") == "no-referrer" and r.headers.get("cache-control") == "no-store"
          and r.headers.get("x-frame-options") == "DENY")
    check("no page can be framed", get(admin, "/").headers.get("x-frame-options") == "DENY")

    # files
    outside = _TMP / "test.sqlite3"
    db.log_attachment(999001, 1, "ISB1", 1, "stolen.db", str(outside), None, False)
    aid = next(a["id"] for a in db.attachments_for_site("ISB1") if a["message_id"] == 999001)
    check("a database row pointing outside the archive can't be used to read other files",
          get(admin, f"/media/file/{aid}").status_code == 404)
    check("a missing file is a 404, not a crash", get(admin, "/media/file/99999999").status_code == 404)

    # accounts
    check("an unknown role is refused",
          post(admin, "/users", {"email": "r@t.test", "name": "R", "password": "longenough1", "role": "root"}).status_code == 400)
    check("a too-short temporary password is refused",
          "8 characters" in msg_of(post(admin, "/users", {"email": "s@t.test", "name": "S", "password": "short", "role": "hq_staff"}))
          and db.get_portal_user_by_email("s@t.test") is None)
    post(admin, "/users", {"email": "fresh@t.test", "name": "Fresh", "password": "temp-password-1", "role": "company_manager", "company_id": str(ids["A"])})
    fresh = anonymous()
    r = fresh.post("/login", data={"email": "fresh@t.test", "password": "temp-password-1"}, follow_redirects=False)
    check("a new account is sent to choose its own password", r.headers["location"] == "/account/password")
    check("...and can't reach the new pages until it has",
          all(get(fresh, p).headers.get("location") == "/account/password" for p in ("/", "/technicians", "/sites", "/inventory")))
    check("the forced-change page itself is reachable", get(fresh, "/account/password").status_code == 200)


# =============================================================================
# 9. EVERY PAGE, EVERY ROLE
# =============================================================================

def test_every_page():
    section("9. Every page renders, for every role")
    hq_pages = ["/", "/sites", "/sites/ISB1", "/sites/ISB1/technical", "/technicians", f"/technicians/{ids['techA']}",
                 "/blockers", "/inventory", "/companies", "/discord", "/hermes/rules", "/users", "/reconcile",
                 "/account/password"]
    manager_pages = ["/", "/sites", "/sites/ISB1", "/technicians", f"/technicians/{ids['techA']}", "/blockers", "/account/password"]
    admin_only = ("/inventory", "/companies", "/discord", "/hermes/rules", "/users", "/reconcile")
    for who, pages in (("admin", hq_pages), ("staff", [p for p in hq_pages if p not in admin_only]), ("mgrA", manager_pages)):
        client = login(who)
        bad = {p: get(client, p).status_code for p in pages if get(client, p).status_code != 200}
        check(f"{who}: {len(pages)} pages all render", not bad, str(bad))
    with db.get_conn() as conn:
        conn.execute("UPDATE portal_users SET company_id = ? WHERE email = 'orphan@t.test'", (db.ensure_company("Empty Co"),))
    empty = login("orphan")
    bad = {p: get(empty, p).status_code for p in ("/", "/sites", "/technicians", "/blockers") if get(empty, p).status_code != 200}
    check("a company with nothing yet: pages render with helpful empty states", not bad, str(bad))
    with db.get_conn() as conn:
        conn.execute("UPDATE portal_users SET company_id = NULL WHERE email = 'orphan@t.test'")
    unlinked = login("orphan")
    check("a company account with no company: dashboard explains, doesn't crash",
          get(unlinked, "/").status_code == 200 and "isn't linked to a company" in get(unlinked, "/").text)
    link = db.ensure_link(technician_id=ids["techA"], created_by="t")
    anon = anonymous()
    check("a technician's link opens without signing in", get(anon, f"/j/{link['token']}").status_code == 200)
    check("a made-up link is a plain 'doesn't work' page (404)",
          get(anon, "/j/not-a-real-token").status_code == 404 and "doesn't work" in get(anon, "/j/not-a-real-token").text)


# =============================================================================
# 10. THE DEPLOY SMOKE HARNESS: works, and can detect a leak
# =============================================================================

def test_smoke_harness():
    import io
    from contextlib import redirect_stdout
    from portal import auth, main
    from tools import smoke_portal

    section("10. The deploy smoke harness works, and can detect a leak")
    real_current_user, real_visible = main.current_user, auth.visible_site_ids

    def run_harness():
        buf = io.StringIO()
        code = 0
        with redirect_stdout(buf):
            try:
                smoke_portal.main_()
            except SystemExit as exc:
                code = int(exc.code or 0)
        main.current_user = real_current_user
        return code, buf.getvalue()

    code, out = run_harness()
    check("healthy data: every page checks out for admin, staff and a company manager",
          code == 0 and "as HQ admin" in out and "as HQ staff" in out and "as company manager" in out, out[-400:])
    check("...including asking a company for another company's site (must be 404)",
          "must be 404" in out and "LEAK" not in out)
    check("...and it does not serve admin pages to a company manager", "must be refused" not in out)
    check("...and it leaves a person's link and Discord's sign-in alone", "/oauth/discord/callback" not in
          [line.split()[-1] for line in out.splitlines() if line.strip().startswith(("200", "30", "40", "50"))])

    leaky = lambda user: None if user["role"] == "company_manager" else real_visible(user)
    auth.visible_site_ids, main.visible_site_ids = leaky, leaky
    code, out = run_harness()
    auth.visible_site_ids = main.visible_site_ids = real_visible
    check("with scoping broken, the harness fails the deploy and names the leak", code == 1 and "LEAK" in out, out[-300:])


# =============================================================================
# 11. SERVERS: which are needed, claiming, releasing, and the admin tools
# =============================================================================

def test_servers():
    import io
    from contextlib import redirect_stdout
    from tools import register_guild, set_secret

    section("11. Servers — which are needed, and the tools around them")
    X = db.ensure_company("Xpress Co")
    db.upsert_sites([site_row(f"XP{i}", "Multan", region="Central B") for i in range(1, 6)], "t")
    db.bulk_assign_company([f"XP{i}" for i in range(1, 6)], X)
    real_size = CONFIG.sites_per_server
    object.__setattr__(CONFIG, "sites_per_server", 2)
    try:
        mine = [s for s in servers.needed_servers() if s["company_id"] == X]
        check("a region that outgrows one server asks for _2 and _3",
              [(s["name"], s["sites"]) for s in mine]
              == [("Xpress Co_Central B", 2), ("Xpress Co_Central B_2", 2), ("Xpress Co_Central B_3", 1)], str(mine))
        db.claim_guild(8801, "whatever", X, "Central B", 1, "admin", status="ready")
        mine = [s for s in servers.needed_servers() if s["company_id"] == X]
        check("a claimed server with room counts, so fewer are asked for",
              [(s["name"], s["sites"]) for s in mine] == [("Xpress Co_Central B_2", 2), ("Xpress Co_Central B_3", 1)], str(mine))
    finally:
        object.__setattr__(CONFIG, "sites_per_server", real_size)
    check("a region is matched whatever its case", db.server_name("Xpress Co", "Central B", 1) == "Xpress Co_Central B")
    before = dict(db.get_guild(8801))
    check("a server that's already a company's can't be claimed over for another (even in a race past the page's checks)",
          db.claim_guild(8801, "Hijack", ids["A"], "North", 2, "admin") is False
          and {k: db.get_guild(8801)[k] for k in ("name", "company_id", "region", "seq", "status")}
          == {k: before[k] for k in ("name", "company_id", "region", "seq", "status")})

    admin = login("admin")
    Y = db.ensure_company("Yonder Co")
    db.upsert_sites([site_row("YD1", "Gilgit", region="North")], "t")
    db.bulk_assign_company(["YD1"], Y)
    page = get(admin, "/discord").text
    check("the Discord page lists each server to create, with its exact name and an Add-the-bot button",
          'value="Yonder Co_North"' in page and f"/discord/add-bot?company_id={Y}&amp;region=North&amp;seq=1" in page)
    check("...and the dashboard says servers are needed",
          "Discord servers needed" in get(admin, "/").text or "Discord server needed" in get(admin, "/").text)
    r = get(admin, f"/discord/add-bot?company_id={X}&region=central b&seq=2")
    target = r.headers.get("location", "")
    check("Add the bot sends the admin to Discord's bot authorisation, asking for a code back to us",
          r.status_code == 302 and target.startswith("https://discord.com/oauth2/authorize?")
          and "scope=bot" in target and "permissions=8" in target and "response_type=code" in target
          and "redirect_uri=https%3A%2F%2Fportal.test%2Foauth%2Fdiscord%2Fcallback" in target
          and "client_id=424242424242" in target)
    check("an unknown company or region is refused",
          get(admin, f"/discord/add-bot?company_id=99999&region=North&seq=1").status_code == 404
          and get(admin, f"/discord/add-bot?company_id={X}&region=Mars&seq=1").status_code == 404)

    # the server the bot was added to by mistake
    db.register_guild_seen(8802, "Personal server", 42)
    fake.reset()
    r = post(admin, "/discord/guilds/8802/leave")
    check("the bot leaves a server it was added to by mistake", fake.left == [8802] and db.get_guild(8802)["status"] == "left")
    check("...but never one that's in use", post(admin, "/discord/guilds/8801/leave").status_code == 404 and not fake.left[1:])
    db.claim_guild(8803, "Old one", X, "Central A", 1, "admin", status="review")
    post(admin, "/discord/guilds/8803/confirm")
    check("confirming a server held for review lets the bot set it up", db.get_guild(8803)["status"] == "confirmed")

    # register_guild: the pilot, which predates claiming
    calls = []

    async def answering(method, path, json=None, reason=None):
        calls.append((method, path, json))
        if method == "GET" and path == "/guilds/8804":
            return {"id": "8804", "name": "RMS PreSurvey"}
        if method == "GET" and path == "/guilds/8804/channels":
            return [{"id": "91", "type": 4, "name": "ISB Sites 001-020"}, {"id": "92", "type": 0, "name": "isb0001"}]
        return {}

    db.upsert_site_channel("PILOT1", "Islamabad", 8804, 92, "ISB Sites 001-020", None, "brief")
    saved_argv, saved_request = sys.argv, discord_api._request
    discord_api._request = answering
    try:
        sys.argv = ["register_guild", "--guild", "8804", "--company", "HQ Tech", "--region", "north", "--seq", "2",
                    "--rename", "HQ Tech_North_2", "--rename-category", "ISB Sites 001-020=Sites 001-050"]
        with redirect_stdout(io.StringIO()):
            register_guild.main()
        row = db.get_guild(8804)
        check("register_guild makes a server set up by hand (the pilot) the company's server for its region, ready",
              row["status"] == "ready" and (row["company_id"], row["region"], row["seq"]) == (ids["A"], "North", 2)
              and row["name"] == "HQ Tech_North_2")
        check("...renames the server and its category in Discord, and the records follow",
              ("PATCH", "/guilds/8804", {"name": "HQ Tech_North_2"}) in calls
              and ("PATCH", "/channels/91", {"name": "Sites 001-050"}) in calls
              and db.get_site_channel("PILOT1")["category_name"] == "Sites 001-050")
        with redirect_stdout(io.StringIO()):
            register_guild.main()
        check("running it again changes nothing", db.get_guild(8804)["status"] == "ready")
        sys.argv = ["register_guild", "--guild", "8805", "--company", "HQ Tech", "--region", "North"]
        try:
            with redirect_stdout(io.StringIO()):
                register_guild.main()
            refused = None
        except SystemExit as exc:
            refused = str(exc)
        check("...and it refuses a slot another server already has", refused and "already" in refused, str(refused))
    finally:
        sys.argv, discord_api._request = saved_argv, saved_request

    # set_secret: into .env without echo
    env = _TMP / "env-test"
    env.write_text("DISCORD_BOT_TOKEN='x'\nSMTP_PASSWORD='old'\n", encoding="utf-8")
    set_secret.set_key(env, "SMTP_PASSWORD", set_secret.clean("abcd efgh ijkl mnop"))
    set_secret.set_key(env, "DISCORD_CLIENT_SECRET", "s3cr3t")
    text = env.read_text(encoding="utf-8")
    check("set_secret replaces a key in place, adds a new one, and drops a Gmail app password's spaces",
          text == "DISCORD_BOT_TOKEN='x'\nSMTP_PASSWORD='abcdefghijklmnop'\nDISCORD_CLIENT_SECRET='s3cr3t'\n", repr(text))
    try:
        set_secret.set_key(env, "SMTP_PASSWORD", "it's\nbroken")
        rejected = False
    except ValueError:
        rejected = True
    check("...and refuses a value that would corrupt the file (quotes, line breaks)",
          rejected and "abcdefghijklmnop" in env.read_text(encoding="utf-8"))
    if os.name == "posix":
        check("...and keeps .env readable by its owner only", (env.stat().st_mode & 0o777) == 0o600)

    # Discord values are checked with Discord before they're saved (faked here). A bot token pasted
    # as the client secret is exactly the mix-up that took the live bot down once.
    import base64
    fake_token = base64.urlsafe_b64encode(b"424242424242").decode().rstrip("=") + ".GAbCdE." + "t" * 38
    calls, answers = [], []

    def fake_discord(method, path, headers, data=None):
        calls.append((method, path))
        return answers.pop(0) if answers else (200, {"id": "424242424242"})
    real_discord = set_secret._discord
    set_secret._discord = fake_discord
    try:
        check("a bot token given as the client secret is refused, without even asking Discord",
              "bot token" in (set_secret.check_with_discord("DISCORD_CLIENT_SECRET", fake_token, "424242424242") or "")
              and not calls)
        check("...and something that isn't a bot token isn't taken as one",
              "isn't a bot token" in (set_secret.check_with_discord("DISCORD_BOT_TOKEN", "s3cr3t-value", "424242424242") or "")
              and not calls)
        answers[:] = [(200, {"id": "424242424242"}), (401, {"message": "401: Unauthorized"}), (200, {"id": "999"})]
        results = [set_secret.check_with_discord("DISCORD_BOT_TOKEN", fake_token, "424242424242") for _ in range(3)]
        check("a bot token is saved only if Discord signs it in as this application's bot",
              results[0] is None and "refused" in (results[1] or "") and "different application" in (results[2] or "")
              and calls[0] == ("GET", "/users/@me"), str(results))
        answers[:] = [(200, {"access_token": "x"}), (401, {"error": "invalid_client"})]
        good = set_secret.check_with_discord("DISCORD_CLIENT_SECRET", "a" * 32, "424242424242")
        bad = set_secret.check_with_discord("DISCORD_CLIENT_SECRET", "b" * 32, "424242424242")
        check("a client secret is saved only if Discord's sign-in accepts it",
              good is None and "invalid_client" in (bad or "") and calls[-1] == ("POST", "/oauth2/token"))
    finally:
        set_secret._discord = real_discord
    env.write_text("DISCORD_CLIENT_ID=424242   # the app\nOTHER='x'\n", encoding="utf-8")
    check("...using the client id as .env has it (with or without quotes and comments)",
          set_secret.env_value(env, "DISCORD_CLIENT_ID") == "424242" and set_secret.env_value(env, "OTHER") == "x"
          and set_secret.env_value(env, "MISSING") is None)


# =============================================================================
# 12. DISCORD'S ANSWERS: the real calls, with the network stubbed
# =============================================================================

class _Resp:
    def __init__(self, status, body):
        self.status, self._body = status, body

    async def json(self, content_type=None):
        return self._body

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False


class _Session:
    """aiohttp.ClientSession, answering from a script of (status, body)."""
    script, seen = [], []

    def __init__(self, *a, **kw):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False

    def request(self, method, url, headers=None, json=None):
        _Session.seen.append((method, url, json))
        return _Resp(*_Session.script.pop(0))

    def post(self, url, data=None, auth=None):
        _Session.seen.append(("POST", url, data))
        return _Resp(*_Session.script.pop(0))

    def get(self, url, headers=None):
        _Session.seen.append(("GET", url, headers))
        return _Resp(*_Session.script.pop(0))


def test_discord_answers():
    import aiohttp
    section("12. Discord's answers — the real calls, with the network stubbed")
    real_session, real_sleep = aiohttp.ClientSession, asyncio.sleep
    saved = {name: getattr(discord_api, name) for name in ("_request", "set_overwrite", "channel_member_overwrites", "leave_guild")}
    for name in saved:
        setattr(discord_api, name, REAL[name])
    aiohttp.ClientSession = _Session

    async def no_wait(seconds):
        _Session.seen.append(("sleep", seconds, None))
    asyncio.sleep = no_wait
    try:
        _Session.script, _Session.seen = [(204, None)], []
        run(discord_api.set_overwrite(11, 22, access.READ_ONLY))
        method, url, body = _Session.seen[0]
        check("a read-only overwrite sends exactly those bits, as a member overwrite",
              method == "PUT" and url.endswith("/channels/11/permissions/22")
              and body == {"allow": "66560", "deny": "1759598709233728", "type": 1})
        _Session.script = [(200, {"permission_overwrites": [
            {"id": "5", "type": 1, "allow": "101440", "deny": "103079215104"},
            {"id": "6", "type": 0, "allow": "1024", "deny": "0"},
            {"id": "7", "type": 1, "allow": "66560", "deny": "1759598709233728"}]})]
        got = run(discord_api.channel_member_overwrites(11))
        check("reading overwrites keeps members only, as numbers", got == {5: (101440, 103079215104), 7: (66560, 1759598709233728)})
        check("...and each reads as its mode",
              access.mode_of(*got[5]) == access.WRITABLE and access.mode_of(*got[7]) == access.READ_ONLY
              and access.mode_of(101440, 0) == access.WRITABLE and access.mode_of(0, 1024) is None)

        _Session.script, _Session.seen = [(429, {"retry_after": 0.25}), (200, {"ok": True})], []
        check("a rate limit is waited out, then the call goes through",
              run(discord_api._request("GET", "/x")) == {"ok": True} and ("sleep", 0.25, None) in _Session.seen)
        _Session.script = [(429, {"retry_after": 0.1})] * 4
        try:
            run(discord_api._request("GET", "/x"))
            gave_up = None
        except discord_api.DiscordError as exc:
            gave_up = exc.status
        check("...but not for ever", gave_up == 429)
        _Session.script = [(404, {"code": 10003, "message": "Unknown Channel"})]
        try:
            run(discord_api._request("GET", "/x"))
            code = None
        except discord_api.DiscordError as exc:
            code = (exc.status, exc.code)
        check("an error carries Discord's status and code", code == (404, 10003))
        _Session.script, _Session.seen = [(204, None)], []
        run(discord_api.leave_guild(55))
        check("leaving a server is DELETE /users/@me/guilds/<id>", _Session.seen[0][:2] == ("DELETE", discord_api.API + "/users/@me/guilds/55"))

        # Discord's sign-in
        _Session.script, _Session.seen = [(200, {"access_token": "A", "refresh_token": "R", "expires_in": 604800,
                                                 "scope": "identify guilds.join"})], []
        answer = run(REAL["oauth.exchange_code"]("the-code"))
        _, url, form = _Session.seen[0]
        check("exchanging a code posts it back with our redirect address",
              answer["access_token"] == "A" and url == oauth.TOKEN_URL and form["grant_type"] == "authorization_code"
              and form["code"] == "the-code" and form["redirect_uri"] == "https://portal.test/oauth/discord/callback")
        _Session.script = [(400, {"error": "invalid_grant"})]
        try:
            run(REAL["oauth.refresh"]("old"))
            kind = None
        except oauth.OAuthError as exc:
            kind = exc.invalid_grant
        check("a refresh Discord refuses is told apart ('connect again')...", kind is True)
        _Session.script = [(503, "unavailable")]
        try:
            run(REAL["oauth.refresh"]("old"))
            kind = None
        except oauth.OAuthError as exc:
            kind = exc.invalid_grant
        check("...from Discord just being down (try again later)", kind is False)
        _Session.script = [(200, {"id": "123", "username": "sultan"})]
        check("asking who signed in", run(REAL["oauth.get_me"]("A"))["username"] == "sultan"
              and _Session.seen[-1][2] == {"Authorization": "Bearer A"})
    finally:
        aiohttp.ClientSession, asyncio.sleep = real_session, real_sleep
        for name, fn in saved.items():
            setattr(discord_api, name, fn)

    perms = discord.Permissions
    check("the two modes' bits are exactly what discord.py means by them",
          access.overwrite_for(access.WRITABLE) == (
              perms(view_channel=True, send_messages=True, add_reactions=True, attach_files=True, read_message_history=True).value,
              perms(create_public_threads=True, create_private_threads=True).value)
          and access.overwrite_for(access.READ_ONLY)[0] == perms(view_channel=True, read_message_history=True).value
          and perms(access.overwrite_for(access.READ_ONLY)[1]).send_messages
          and perms(access.overwrite_for(access.READ_ONLY)[1]).add_reactions)


# =============================================================================
# 13. ALIGNMENT: every number sits under its heading, on every page
# =============================================================================

def test_alignment():
    from tools.smoke_portal import STYLE_CSS, css_right_aligns_numeric_headers, misaligned_columns

    section("13. Alignment — headings line up with their numbers, on every page")
    hq_pages = ["/", "/sites", "/sites?show=all", "/sites/ISB1", "/sites/ISB1/technical", "/technicians",
                 f"/technicians/{ids['techA']}", "/blockers", "/inventory", "/companies", "/discord", "/hermes/rules",
                 "/users", "/reconcile"]
    manager_pages = ["/", "/sites", "/sites/ISB1", "/technicians", f"/technicians/{ids['techA']}", "/blockers"]
    for who, pages in (("admin", hq_pages), ("staff", [p for p in hq_pages if p not in
                       ("/inventory", "/companies", "/discord", "/hermes/rules", "/users", "/reconcile")]), ("mgrA", manager_pages)):
        client = login(who)
        bad = {}
        for page in pages:
            r = get(client, page)
            problems = misaligned_columns(r.text) if r.status_code == 200 else [f"status {r.status_code}"]
            if problems:
                bad[page] = problems
        check(f"{who}: every table on {len(pages)} pages has headings aligned like their cells", not bad, str(bad)[:600])

    check("the stylesheet right-aligns numeric headings, not just the numbers",
          css_right_aligns_numeric_headers(STYLE_CSS.read_text(encoding="utf-8")))
    check("...and the check notices when it doesn't (the original fault)",
          not css_right_aligns_numeric_headers("th, td { text-align: left; }\ntd.num { text-align: right; }"))
    broken = ("<table><tr><th>Site</th><th class='num'>Techs</th></tr>"
              "<tr><td>ISB1</td><td>3</td></tr></table>")
    check("the table check catches a numeric heading over a left-aligned cell", misaligned_columns(broken))
    check("...and passes a table whose columns agree",
          not misaligned_columns("<table><tr><th>Site</th><th class='num'>Techs</th></tr>"
                                 "<tr><td>ISB1</td><td class='num'>3</td></tr>"
                                 "<tr><td colspan='2'>nothing else</td></tr></table>"))


# =============================================================================
# 14. COMPANIES: only an empty one can be removed
# =============================================================================

def test_company_removal():
    section("14. Companies — only an empty one can be removed")
    admin = login("admin")

    def company_id(name):
        company = db.get_company_by_name(name)
        return company["id"] if company else None

    post(admin, "/companies", {"name": "Removable Co"})
    gone = company_id("Removable Co")
    r = post(admin, f"/companies/{gone}/remove")
    check("an empty company is removed", r.status_code == 303 and company_id("Removable Co") is None
          and "Removed Removable Co" in msg_of(r), msg_of(r))
    post(admin, "/companies", {"name": "Removable Co"})
    check("...and its name can be used again", company_id("Removable Co") is not None)

    busy = company_id("Removable Co")
    db.upsert_sites([site_row("RMV1", "Lahore")], "t")
    db.bulk_assign_company(["RMV1"], busy)
    tech = db.create_technician(busy, "Rem Tech", "0300 7770001", None, "t")
    db.create_portal_user("rem@t.test", hash_password(PASSWORD), "Rem", "company_manager",
                          must_change=False, company_id=busy)
    r = post(admin, f"/companies/{busy}/remove")
    check("a company with a site, a technician and a login is refused, and told why",
          company_id("Removable Co") == busy and "1 site, 1 technician, 1 login" in msg_of(r), msg_of(r))

    db.bulk_assign_company(["RMV1"], None)
    db.remove_technician(tech)
    with db.get_conn() as conn:
        conn.execute("DELETE FROM portal_users WHERE email = 'rem@t.test'")
    r = post(admin, f"/companies/{busy}/remove")
    check("once they're taken back it can go; a removed technician's history doesn't block it",
          company_id("Removable Co") is None, msg_of(r))

    post(admin, "/companies", {"name": "Server Co"})
    with_server = company_id("Server Co")
    db.claim_guild(8850, "Server Co_North", with_server, "North", 1, "admin", status="ready")
    r = post(admin, f"/companies/{with_server}/remove")
    check("a company whose Discord server still exists is refused, even with no sites in it, and told what to do",
          company_id("Server Co") == with_server and "1 Discord server" in msg_of(r) and "Release" in msg_of(r), msg_of(r))
    db.mark_guild_left(8850)
    r = post(admin, f"/companies/{with_server}/remove")
    check("...once the server is gone from Discord it can go, and the Discord page still works",
          company_id("Server Co") is None and get(admin, "/discord").status_code == 200, msg_of(r))

    check("a company with sites (HQ Tech) is never removed",
          post(admin, f"/companies/{ids['A']}/remove").status_code == 303 and db.get_company(ids["A"]) is not None)
    check("an unknown company is a 404", post(admin, "/companies/999999/remove").status_code == 404)
    post(admin, "/companies", {"name": "Guarded Co"})
    guarded = company_id("Guarded Co")
    check("HQ staff and company managers can't remove companies",
          post(login("staff"), f"/companies/{guarded}/remove").status_code == 403
          and post(login("mgrA"), f"/companies/{guarded}/remove").status_code == 403
          and company_id("Guarded Co") == guarded)
    page = get(admin, "/companies").text
    check("the Remove button is offered only for companies with nothing in them",
          f"/companies/{guarded}/remove" in page and f"/companies/{ids['A']}/remove" not in page)


# =============================================================================
# 15. LIVE PAGES: they refresh on change, and only for what the viewer may see
# =============================================================================

def test_live():
    import time as _time
    from itertools import count

    section("15. Live pages — they refresh on change, and only for what the viewer may see")
    admin, mgr_a, mgr_b = login("admin"), login("mgrA"), login("mgrB")
    version = lambda client: get(client, "/live/version").json()["v"]
    message_ids = count(880_000)

    first = version(mgr_a)
    check("a company's version holds still while nothing changes", version(mgr_a) == first)
    page = get(mgr_a, "/").text
    check("a page carries the version it was drawn at, and loads the live script",
          f'data-live-version="{first}"' in page and "/static/live.js" in page)
    check("the version isn't cached anywhere on the way", get(mgr_a, "/live/version").headers.get("cache-control") == "no-store")

    def scoped_ok():
        a0, b0, p0 = version(mgr_a), version(mgr_b), version(admin)
        db.log_message(next(message_ids), CH["KHI1"], "KHI1", 2, "Bravo Tech", "news on B's site", _time.time())
        return version(mgr_a) == a0 and version(mgr_b) != b0 and version(admin) != p0

    check("activity on another company's site leaves this company's version alone (theirs and HQ's move)", scoped_ok())
    a0 = version(mgr_a)
    blocker = db.open_blocker("ISB2", CH["ISB2"], next(message_ids), 1, "Alpha Tech", "live: no power", "no power")
    check("a new blocker on the company's own site moves its version", version(mgr_a) != a0)
    a1 = version(mgr_a)
    db.resolve_blocker(blocker, "t", "resolved")
    check("...and so does resolving it", version(mgr_a) != a1)
    b0 = version(mgr_b)
    db.set_site_progress("ISB2", "on_site", "t")
    check("progress on one company's site doesn't touch another's version", version(mgr_b) == b0)
    check("signed out, the version isn't served", get(TestClient(app, base_url="https://testserver"), "/live/version").status_code == 303)
    check("a company login with no company always sees the same (empty) version",
          version(login("orphan")) == "none")

    duplicated = {}
    for who, pages in (("admin", ["/", "/sites", "/sites/ISB1", "/technicians", f"/technicians/{ids['techA']}",
                                  "/blockers", "/inventory", "/discord"]),
                       ("mgrA", ["/", "/sites", "/sites/ISB1", "/technicians", "/blockers"])):
        client = login(who)
        for p in pages:
            found = re.findall(r'data-live\s+id="([^"]+)"', get(client, p).text)
            if not found or len(found) != len(set(found)):
                duplicated[f"{who} {p}"] = found
    check("every live page has live regions, each with its own id", not duplicated, str(duplicated))

    real = db.change_fingerprint
    db.change_fingerprint = lambda company_id=None: real(None)       # a fingerprint that ignores the company
    try:
        leaky = scoped_ok()
    finally:
        db.change_fingerprint = real
    check("with the scoping broken on purpose, the check above goes red", leaky is False)


# =============================================================================
# 16. CONNECTING DISCORD: a person's link, and adding the bot to a server
# =============================================================================

def test_connect():
    from urllib.parse import parse_qs, urlparse
    from portal import main
    section("16. Connecting Discord — one link per person, and the bot into new servers")
    signin.calls.clear()
    anon = anonymous()

    tech = fresh_tech(ids["A"], "Imran Shah")
    assign.assign_sites(tech, ["ISB2"], "mgr")
    link = db.current_link(technician_id=tech["id"])
    page = get(anon, f"/j/{link['token']}").text
    check("their link greets them, names the company and their sites, and offers one button",
          "Imran" in page and "HQ Tech" in page and "ISB2" in page and f'href="/j/{link["token"]}/go"' in page
          and "Connect Discord" in page)
    r = get(anon, f"/j/{link['token']}/go")
    target = urlparse(r.headers.get("location", ""))
    q = parse_qs(target.query)
    check("the button goes to Discord's own sign-in, asking to know who they are and to add them to servers",
          r.status_code == 302 and target.netloc == "discord.com" and q["scope"] == ["identify guilds.join"]
          and q["response_type"] == ["code"] and q["redirect_uri"] == ["https://portal.test/oauth/discord/callback"])
    state = q["state"][0]
    check("the note that travels with it names the link by number, never its token",
          link["token"] not in r.headers["location"] and main._signer.loads(state)["link_id"] == link["id"])

    def answer(code, state_, **extra):
        return get(anon, "/oauth/discord/callback?" + "&".join(
            f"{k}={v}" for k, v in {"code": code, "state": state_, **extra}.items() if v is not None))

    signin.person("c-imran", 6601, "imran.shah")
    r = answer("c-imran", state)
    tech = db.get_technician(tech["id"])
    check("coming back from Discord connects them and sends them to their page",
          r.status_code == 303 and r.headers["location"] == f"/j/{link['token']}" and tech["discord_id"] == 6601)
    check("...keeps what the bot needs to add them to servers later",
          db.get_tokens(6601)["refresh_token"].startswith("ref-6601") and db.get_tokens(6601)["username"] == "imran.shah")
    page = get(anon, f"/j/{link['token']}").text
    check("their page now says they're connected, and their site is on its way",
          "Connected" in page and "@imran.shah" in page and "Opening now" in page and 'http-equiv="refresh"' in page)
    db.assign_tech("ISB2", 6601, "bot")
    db.finish_queue_row(db.queued_row(tech["id"], "ISB2")["id"])
    page = get(anon, f"/j/{link['token']}").text
    check("once the bot has opened it, the page has a button straight into the channel",
          f"https://discord.com/channels/1001/{CH['ISB2']}" in page)

    # --- what's refused on the way back
    before = signin.calls["exchange"]
    check("an answer with a forged note is refused, and Discord is never asked",
          answer("anything", "forged.note.here").status_code == 400 and signin.calls["exchange"] == before)
    old = URLSafeState.expired({"kind": "join", "link_id": link["id"], "nonce": "x"})
    check("...and so is one whose note has expired", answer("anything", old).status_code == 400 and signin.calls["exchange"] == before)
    check("...and one with no note at all", get(anon, "/oauth/discord/callback?code=x").status_code == 400)
    other = fresh_tech(ids["A"], "Other One")
    other_link = db.ensure_link(technician_id=other["id"], created_by="t")
    other_state = main._signer.dumps({"kind": "join", "link_id": other_link["id"], "nonce": "n"})
    r = answer(None, other_state, error="access_denied")
    check("pressing Cancel in Discord brings them back to try again, nothing connected",
          r.status_code == 200 and "Not connected yet" in r.text and db.get_technician(other["id"])["discord_id"] is None)
    signin.person("c-narrow", 6602, "narrow", scope="identify")
    r = answer("c-narrow", other_state)
    check("a sign-in without permission to add them to servers isn't accepted",
          "permission" in r.text and db.get_technician(other["id"])["discord_id"] is None)
    signin.person("c-taken", 6601, "imran.shah")                     # Imran's account, on Other One's link
    r = answer("c-taken", other_state)
    check("an account already connected for someone else is refused",
          r.status_code == 409 and "someone else" in r.text and db.get_technician(other["id"])["discord_id"] is None)
    signin.person("c-second", 6603, "second.account")                # a different account on Imran's link
    r = answer("c-second", main._signer.dumps({"kind": "join", "link_id": link["id"], "nonce": "n"}))
    check("Imran's link can't be used by another account once he's connected",
          r.status_code == 409 and db.get_technician(tech["id"])["discord_id"] == 6601)
    run(assign.remove_technician(db.get_technician(other["id"])))
    check("a removed technician's link doesn't open", get(anon, f"/j/{other_link['token']}").status_code == 404)
    with db.get_conn() as conn:
        conn.execute("UPDATE join_links SET expires_at = ? WHERE technician_id = ?", (time.time() - 1, ids["techA"]))
    stale = db.current_link(technician_id=ids["techA"])
    if stale:
        page = get(anon, f"/j/{stale['token']}")
        check("an expired, never-used link says so, and its button doesn't go to Discord",
              "has expired" in page.text and get(anon, f"/j/{stale['token']}/go").headers["location"] == f"/j/{stale['token']}")

    # --- HQ staff
    admin = login("admin")
    r = post(admin, "/discord/staff", {"name": "Bilal Engineer", "discord_role": "Engineer", "email": "bilal@t.test",
                                       "phone": "0300 7654321", "country": "PK"})
    bilal = next(s for s in db.all_staff() if s["name"] == "Bilal Engineer")
    staff_link = db.current_link(staff_id=bilal["id"])
    check("adding HQ staff makes their link and emails it", staff_link and "on its way to bilal@t.test" in msg_of(r)
          and bilal["phone"] == "+923007654321")
    page = get(anonymous(), f"/j/{staff_link['token']}").text
    check("...their link says what it's for", "HQ's team" in page and "Engineer" in page)
    signin.person("c-bilal", 6701, "bilal.e")
    get(anonymous(), "/oauth/discord/callback?code=c-bilal&state=" +
        main._signer.dumps({"kind": "join", "link_id": staff_link["id"], "nonce": "n"}))
    check("connecting links the staff member (the bot then puts them into every server)",
          db.get_staff(bilal["id"])["discord_id"] == 6701 and "Connected" in get(admin, "/discord").text)
    r = post(admin, "/discord/staff/me")
    me = db.staff_by_email("admin@t.test")
    check("'Connect my own Discord' makes the admin Management staff and goes straight to Discord",
          me and me["discord_role"] == "Management" and r.headers["location"].endswith("/go"))

    # --- adding the bot to a new server
    Z = db.ensure_company("Zeta Co")
    signin.calls.clear()

    def bot_state(uid, seq=1, region="North", company=Z):
        return main._signer.dumps({"kind": "bot", "company_id": company, "region": region, "seq": seq,
                                   "uid": uid, "nonce": "n"})

    signin.bot("b-ok", 9901, "Zeta server")
    r = get(admin, f"/oauth/discord/callback?code=b-ok&state={bot_state(ids['admin'])}&guild_id=9901")
    row = db.get_guild(9901)
    check("after the admin adds the bot, the server is claimed for that company and region",
          row and row["status"] == "claimed" and (row["company_id"], row["region"], row["seq"]) == (Z, "North", 1)
          and "being set up" in msg_of(r))
    signin.bot("b-dup", 9902)
    r = get(admin, f"/oauth/discord/callback?code=b-dup&state={bot_state(ids['admin'])}&guild_id=9902")
    check("a slot that already has a server is refused before Discord is even asked",
          "already has a server" in msg_of(r) and db.get_guild(9902) is None and signin.calls["exchange"] == 1)
    signin.bot("b-used", 1001)
    r = get(admin, f"/oauth/discord/callback?code=b-used&state={bot_state(ids['admin'], seq=2)}&guild_id=1001")
    check("picking a server that's already a site server is refused",
          "already HQ Tech_North" in msg_of(r) and db.get_guild(1001)["company_id"] == ids["A"])
    signin.bot("b-mismatch", 9904)
    r = get(admin, f"/oauth/discord/callback?code=b-mismatch&state={bot_state(ids['admin'], seq=3)}&guild_id=9903")
    check("if Discord's answer names a different server, nothing is claimed",
          "different server" in msg_of(r) and db.get_guild(9903) is None and db.get_guild(9904) is None)
    signin.bot("b-other", 9905)
    r = get(login("staff"), f"/oauth/discord/callback?code=b-other&state={bot_state(ids['admin'], seq=4)}&guild_id=9905")
    check("the answer only counts for the admin who started it (anyone else is sent to sign in)",
          r.headers["location"] == "/login" and db.get_guild(9905) is None)
    r = get(admin, f"/oauth/discord/callback?state={bot_state(ids['admin'], seq=5)}&guild_id=9906&error=access_denied")
    check("cancelling in Discord changes nothing", "cancelled" in msg_of(r) and db.get_guild(9906) is None)


class URLSafeState:
    @staticmethod
    def expired(payload):
        """A note signed long enough ago to have expired."""
        from itsdangerous import URLSafeTimedSerializer
        from itsdangerous.timed import TimestampSigner

        class Old(TimestampSigner):
            def get_timestamp(self):
                return int(time.time()) - 3 * 3600
        return URLSafeTimedSerializer("test-secret", salt="discord-oauth", signer=Old).dumps(payload)


# =============================================================================
# 17. COMPLETING A SURVEY: technicians keep the channel, read-only
# =============================================================================

def test_completion():
    section("17. Completing a survey — technicians can read, not post; and reopen")
    fake.reset()
    mgrA, staff = login("mgrA"), login("staff")
    for uid in (7701, 7702):
        db.assign_tech("ISB2", uid, "t")
        fake.overwrites.setdefault(CH["ISB2"], {})[uid] = access.overwrite_for(access.WRITABLE)

    r = post(mgrA, "/sites/ISB2/complete")
    check("a company manager marks their site's survey complete", db.site_mode("ISB2") == access.READ_ONLY
          and "can still read" in msg_of(r), msg_of(r))
    check("...and every technician on it becomes read-only in Discord, at once",
          all(fake.overwrites[CH["ISB2"]][u] == access.overwrite_for(access.READ_ONLY) for u in (7701, 7702)))
    page = get(mgrA, "/sites/ISB2").text
    check("the site page says so, and offers Reopen", "Complete" in page and "/sites/ISB2/reopen" in page
          and "/sites/ISB2/complete" not in page)
    check("the status can't be changed behind a completed survey's back",
          "Reopen it first" in msg_of(post(mgrA, "/sites/ISB2/progress", {"state": "on_site"}))
          and db.site_mode("ISB2") == access.READ_ONLY)
    check("nobody new can be assigned to it", "complete" in msg_of(post(mgrA, "/sites/ISB2/assign", {"technician_id": str(ids["techA"])})))

    r = post(mgrA, "/sites/ISB2/reopen")
    check("reopening lets them post again", db.site_mode("ISB2") == access.WRITABLE
          and all(fake.overwrites[CH["ISB2"]][u] == access.overwrite_for(access.WRITABLE) for u in (7701, 7702)))
    check("'Done' isn't a status you can just pick (it's the button)",
          "Mark survey complete" in msg_of(post(mgrA, "/sites/ISB2/progress", {"state": "done"}))
          and db.site_mode("ISB2") == access.WRITABLE)
    seen = get(mgrA, "/live/version").json()["v"]
    check("HQ staff can complete a survey too", post(staff, "/sites/ISB2/complete").status_code == 303
          and db.site_mode("ISB2") == access.READ_ONLY)
    page = get(mgrA, "/sites/ISB2").text
    check("...and the company's manager sees it straight away: their open pages refresh, naming who completed it",
          get(mgrA, "/live/version").json()["v"] != seen and "Marked complete by staff" in page
          and "/sites/ISB2/reopen" in page)
    r = post(mgrA, "/sites/ISB2/reopen")
    check("...and can reopen it themselves if it was a mistake",
          r.status_code == 303 and db.site_mode("ISB2") == access.WRITABLE
          and all(fake.overwrites[CH["ISB2"]][u] == access.overwrite_for(access.WRITABLE) for u in (7701, 7702)))

    fake.fail.add("set_overwrite")
    r = post(mgrA, "/sites/ISB2/complete")
    check("if Discord refuses, the page says so and pressing again retries",
          "Press 'Mark survey complete' again" in msg_of(r) and db.site_mode("ISB2") == access.READ_ONLY)
    fake.fail.clear()
    post(mgrA, "/sites/ISB2/complete")
    check("...which then goes through", all(fake.overwrites[CH["ISB2"]][u] == access.overwrite_for(access.READ_ONLY) for u in (7701, 7702)))

    fake.overwrites[CH["ISB2"]][7701] = access.overwrite_for(access.WRITABLE)          # drifted
    admin = login("admin")
    page = get(admin, "/reconcile").text
    check("reconcile spots someone who can still post on a completed site", "7701" in page and "should be read-only" in page)
    post(admin, "/reconcile/apply")
    check("...and fixes it", fake.overwrites[CH["ISB2"]][7701] == access.overwrite_for(access.READ_ONLY))
    post(mgrA, "/sites/ISB2/reopen")
    for uid in (7701, 7702):
        db.revoke_tech("ISB2", uid)
        fake.overwrites[CH["ISB2"]].pop(uid, None)


# =============================================================================
# 18. EMAIL AND WHATSAPP: how the link reaches them
# =============================================================================

def test_messages():
    from urllib.parse import unquote
    section("18. Email and WhatsApp — how the link reaches them")
    person = {"name": "Sultan Shb", "phone": "+16475550123"}
    text = links.whatsapp_text(person, "FieldCo", ["KHI0023", "KHI0024"], "TOKEN123", connected=False)
    check("the WhatsApp message greets them, names the company and the sites, and carries the link",
          text.startswith("Assalam o Alaikum Sultan") and "FieldCo" in text and "KHI0023, KHI0024" in text
          and "https://portal.test/j/TOKEN123" in text and "Connect Discord" in text)
    check("...once connected, it says the sites open by themselves",
          "nayi sites" in links.whatsapp_text(person, "FieldCo", ["KHI0025"], "T", connected=True))
    url = links.whatsapp_url(person["phone"], text)
    check("wa.me gets the number as digits and the message fully encoded",
          url.startswith("https://wa.me/16475550123?text=") and " " not in url and "\n" not in url
          and unquote(url.split("text=", 1)[1]) == text)
    many = links.whatsapp_text(person, "FieldCo", [f"S{i}" for i in range(30)], "T", connected=False)
    check("a long list of sites is cut short, not dumped", "and 18 more" in many)

    subject, body = links.email_for({"name": "Sultan"}, "Evil\r\nBcc: victim@x", ["S1"], "T", connected=False)
    check("an email subject can't be split into extra headers", "\r" not in subject and "\n" not in subject)
    mail.reset()
    subject, body = links.email_for({"name": "Sultan"}, "FieldCo", ["S1"], "T", connected=False)
    email_id = db.log_email("sultan@example.test", subject, "t")
    run(links.deliver(email_id, "sultan@example.test", subject, body))
    sent = mail.sent[0]
    check("mail goes out over STARTTLS, signed in, from the portal's address",
          mail.tls == 1 and mail.logins == ["rms@portal.test"] and sent["From"] == "RMS <rms@portal.test>"
          and mail.connections == [("smtp.test", 587)])
    with db.get_conn() as conn:
        status = conn.execute("SELECT status FROM emails WHERE id = ?", (email_id,)).fetchone()[0]
    check("the email log records it", status == "sent")

    real_config = links.CONFIG
    links.CONFIG = types.SimpleNamespace(email_ready=lambda: False)
    try:
        class Background:
            tasks = []

            def add_task(self, *a):
                self.tasks.append(a)
        background = Background()
        queued = links.queue_email(background, "x@y.test", "s", "b", "t")
    finally:
        links.CONFIG = real_config
    check("with no mail server set up nothing is queued (WhatsApp still works)", queued is None and not background.tasks)
    subject, body = links.staff_email({"name": "Bilal Ali", "discord_role": "Engineer"}, "TOK")
    check("the staff email says what it's for", "Engineer" in body and "https://portal.test/j/TOK" in body)


if __name__ == "__main__":
    install_fakes()
    build_fixtures()
    try:
        test_tenancy()
        test_assignment()
        test_bot_worker()
        test_inventory()
        test_insights()
        test_rules()
        test_database()
        test_security()
        test_every_page()
        test_smoke_harness()
        test_servers()
        test_discord_answers()
        test_alignment()
        test_company_removal()
        test_live()
        test_connect()
        test_completion()
        test_messages()
    finally:
        pass
    print(f"\n{sum(results)}/{len(results)} checks passed"
          + (f" ({len(skipped)} skipped: {NO_REFERENCE})" if skipped else ""))
    shutil.rmtree(_TMP, ignore_errors=True)
    sys.exit(0 if all(results) else 1)
