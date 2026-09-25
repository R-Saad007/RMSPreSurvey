"""The site template: the one layout every site upload uses.

An HQ admin downloads the template, fills in a row per site and uploads it
for a company (only HQ adds sites: a site's company and region decide which
Discord server gets built). Everyone uses the same columns, so an upload
always means the same thing and nobody's spreadsheet habits leak into the
inventory. The customer's own workbook is not accepted here — the CLI
(tools/build_site_csv.py --template) turns it into this template.

Parsing is all-or-nothing: if any row is wrong, nothing is imported and the
first few problems are named by row number, so a half-imported file never
has to be untangled. Separated from the route so it can be tested without a
web server.
"""
import csv
import io
import re
from collections import Counter
from datetime import date, datetime

from storage.db import SITE_COLUMNS

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 2000
MAX_CELL = 1000

# Site ids end up in URLs and channel names. Anything outside this is either a
# typo or a stray row, and would break routing rather than fail politely.
_SITE_ID = re.compile(r"^[A-Za-z0-9_-]{2,32}$")

# (heading, field, required). The headings are what people see; the order is
# the order of the sheet. Planned install date isn't a SITE_COLUMN: it's the
# plan, not the customer's description of the site.
TEMPLATE_COLUMNS = (
    ("Site ID", "site_id", True),
    ("Region", "region", True),
    ("City", "city", True),
    ("District", "district", False),
    ("Priority", "priority", False),
    ("RMS category", "rms_category", False),
    ("Site type", "site_type", False),
    ("Grid", "grid", False),
    ("Power source", "power_source", False),
    ("Tenants", "tenants", False),
    ("Tenancy summary", "tenancy_summary", False),
    ("SIM 1", "sim1", False),
    ("SIM 2", "sim2", False),
    ("Existing equipment", "existing_equipment", False),
    ("Install BOQ", "install_boq", False),
    ("BOQ remarks", "boq_remarks", False),
    ("Smart integration", "smart_integration", False),
    ("Google Maps link", "maps_url", False),
    ("MBU", "mbu", False),
    ("MBU lead", "mbu_lead", False),
    ("MBU lead contact", "mbu_lead_contact", False),
    ("Zonal manager", "zonal_manager", False),
    ("Zonal manager contact", "zonal_manager_contact", False),
    ("Planned install date (YYYY-MM-DD)", "target_install_date", False),
)
assert {f for _, f, _ in TEMPLATE_COLUMNS} == {"site_id", "target_install_date", *SITE_COLUMNS}

SHEET = "Sites"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class UploadError(Exception):
    """The file can't be used; the message says why, for the person uploading."""


def _heading_key(text) -> str:
    """How a heading is matched: case, spacing, the required-marker star and
    the bracketed hint don't matter — the words do."""
    text = re.sub(r"\(.*?\)", "", str(text or ""))
    return " ".join(text.replace("*", " ").split()).lower()


_BY_HEADING = {_heading_key(h): field for h, field, _ in TEMPLATE_COLUMNS}


def _label(heading: str, required: bool) -> str:
    return f"{heading} *" if required else heading


# --- building the template -----------------------------------------------

def build_template(regions, rows=()) -> bytes:
    """The .xlsx people fill in. `rows` pre-fills it (the CLI uses that to
    turn the customer's workbook into this layout)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([_label(h, req) for h, _, req in TEMPLATE_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1C2440")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    widths = {"site_id": 12, "region": 12, "city": 16, "existing_equipment": 40, "install_boq": 40,
              "maps_url": 34, "target_install_date": 30}
    for index, (heading, field, _) in enumerate(TEMPLATE_COLUMNS, start=1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = widths.get(field, max(14, len(heading) + 4))

    last_row = max(MAX_ROWS + 1, len(rows) + 1)
    site_col = get_column_letter(1)
    date_col = get_column_letter(len(TEMPLATE_COLUMNS))
    for r in range(2, last_row + 1):
        ws[f"{site_col}{r}"].number_format = "@"          # keep 0013 as 0013, never 13
        ws[f"{date_col}{r}"].number_format = "yyyy-mm-dd"

    for r, row in enumerate(rows, start=2):
        for c, (_, field, _) in enumerate(TEMPLATE_COLUMNS, start=1):
            value = row.get(field)
            if value in (None, ""):
                continue
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"      # text, never a formula

    lists = wb.create_sheet("Lists")
    for i, name in enumerate(regions, start=1):
        lists.cell(row=i, column=1, value=name)
    lists.sheet_state = "hidden"
    if regions:
        region_col = get_column_letter(2)
        # No leading '=': the file stores the formula text the way Excel itself does.
        pick = DataValidation(type="list", formula1=f"Lists!$A$1:$A${len(regions)}", allow_blank=True,
                              showErrorMessage=True, errorTitle="Unknown region",
                              error="Pick a region from the list. Ask HQ if yours is missing.")
        ws.add_data_validation(pick)
        pick.add(f"{region_col}2:{region_col}{last_row}")

    notes = wb.create_sheet("How to fill")
    for line in (
        "One row per site. Columns marked * are required: Site ID, Region, City.",
        f"Region must be one of: {', '.join(regions) or '(ask HQ)'} — pick it from the dropdown.",
        "Planned install date is optional; type it as YYYY-MM-DD.",
        "Don't rename, add or remove columns — the portal checks them.",
        f"Up to {MAX_ROWS} sites per file. Uploading the same site again updates it; a blank cell keeps what's there.",
        "If anything in the file is wrong, nothing is imported and the portal tells you which rows to fix.",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- reading an upload ---------------------------------------------------

def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _xlsx_rows(data: bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise UploadError("That file couldn't be read as an Excel workbook. "
                          "Download the template from the portal and fill that in.") from exc
    try:
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        for count, row in enumerate(ws.iter_rows(values_only=True)):
            if count > MAX_ROWS + 1:
                raise UploadError(f"That file has more than {MAX_ROWS} sites. Split it into smaller files.")
            yield list(row)
    finally:
        wb.close()


def _csv_rows(data: bytes):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UploadError("That CSV isn't UTF-8 text.") from exc
    for count, row in enumerate(csv.reader(io.StringIO(text))):
        if count > MAX_ROWS + 1:
            raise UploadError(f"That file has more than {MAX_ROWS} sites. Split it into smaller files.")
        yield row


def _parse_day(text: str) -> str:
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError:
        return datetime.fromisoformat(text).date().isoformat()     # '2026-10-01 00:00:00' from a CSV export


def parse_sites_file(filename: str, data: bytes, regions) -> list[dict]:
    """The rows of a filled-in template: dicts of site_id, every SITE_COLUMN
    and target_install_date (ISO text or None). Raises UploadError, naming the
    first few problems, if the file can't be imported as a whole."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadError("That file is over 10 MB — is it the right one?")
    if not data:
        raise UploadError("That file is empty.")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        raw = _xlsx_rows(data)
    elif name.endswith(".csv"):
        raw = _csv_rows(data)
    else:
        raise UploadError("Upload the site template (.xlsx), filled in — download it from this page.")

    raw = iter(raw)
    header = next(raw, None)
    if not header:
        raise UploadError("That file is empty.")
    fields = [_BY_HEADING.get(_heading_key(h)) if _heading_key(h) else None for h in header]
    unknown = [str(h).strip() for h, f in zip(header, fields) if f is None and _heading_key(h)]
    missing = [h for h, f, _ in TEMPLATE_COLUMNS if f not in fields]
    if unknown or missing:
        detail = []
        if missing:
            detail.append("missing " + ", ".join(missing[:4]) + ("…" if len(missing) > 4 else ""))
        if unknown:
            detail.append("unexpected " + ", ".join(unknown[:4]) + ("…" if len(unknown) > 4 else ""))
        raise UploadError("That isn't the site template (" + "; ".join(detail) + "). "
                          "Download the template from this page and fill that in.")

    canonical_region = {r.lower(): r for r in regions}
    rows, problems = [], []
    for number, values in enumerate(raw, start=2):
        texts = [_cell_text(v) for v in values]
        if not any(texts):
            continue
        if len(rows) >= MAX_ROWS:
            raise UploadError(f"That file has more than {MAX_ROWS} sites. Split it into smaller files.")
        row = {field: "" for _, field, _ in TEMPLATE_COLUMNS}
        for field, text in zip(fields, texts):
            if field:
                row[field] = text

        site_id = row["site_id"].upper()
        row["site_id"] = site_id
        if not site_id:
            problems.append(f"row {number}: no Site ID")
        elif not _SITE_ID.match(site_id):
            problems.append(f"row {number}: '{site_id[:20]}' isn't a usable Site ID")
        if not row["region"]:
            problems.append(f"row {number}: no Region")
        elif row["region"].lower() not in canonical_region:
            problems.append(f"row {number}: Region '{row['region'][:20]}' isn't one of {', '.join(regions)}")
        else:
            row["region"] = canonical_region[row["region"].lower()]
        if not row["city"]:
            problems.append(f"row {number}: no City")
        if row["target_install_date"]:
            try:
                row["target_install_date"] = _parse_day(row["target_install_date"])
            except ValueError:
                problems.append(f"row {number}: planned date '{row['target_install_date'][:20]}' isn't YYYY-MM-DD")
        else:
            row["target_install_date"] = None
        long = [h for h, f, _ in TEMPLATE_COLUMNS if len(row[f] or "") > MAX_CELL]
        if long:
            problems.append(f"row {number}: {long[0]} is over {MAX_CELL} characters")
        rows.append(row)

    duplicates = [s for s, n in Counter(r["site_id"] for r in rows if r["site_id"]).items() if n > 1]
    if duplicates:
        problems.append("listed more than once: " + ", ".join(duplicates[:5]) + ("…" if len(duplicates) > 5 else ""))
    if problems:
        shown = "; ".join(problems[:4]) + (f"; and {len(problems) - 4} more" if len(problems) > 4 else "")
        raise UploadError(f"Nothing was imported — fix these and upload again: {shown}.")
    if not rows:
        raise UploadError("No sites were found in that file.")
    return rows
